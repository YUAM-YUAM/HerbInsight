# -*- coding: utf-8 -*-
"""
HerbInsight —— 中药复方多组学机制分析平台
基于网络药理学与多组学数据融合的中药复方作用机制智能分析平台

运行方式：
    cd C:\\Users\\32330\\Desktop\\HerbInsight
    .\\.venv\\Scripts\\Activate.ps1
    streamlit run app.py

浏览器打开 http://localhost:8501
"""

import os
from io import BytesIO
import yaml
import pandas as pd
import streamlit as st


from src import project_manager
from src.enrichment_engine import (
    example_files,
    extract_genes,
    read_table,
    run_over_representation,
    split_gene_text,
)
from src.network_engine import analyze_network, example_edge_file, read_edge_table, map_edge_columns
from src.learning_paths import page_learning_paths
from src.navigation import NAV_GROUPS, ROUTES, resolve_route, course_neighbors

# ---------------------------------------------------------------------------
# 全局配置
# ---------------------------------------------------------------------------
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECTS_DIR = os.path.join(BASE_DIR, "projects")

PAGE_CONFIG = {
    "page_title": "HerbInsight · 网络药理学教学实训",
    "page_icon": "🌿",
    "layout": "wide",
}

DISCLAIMER = "本系统用于科研辅助和教育展示，不用于临床诊断、治疗决策或未经验证的疗效宣传。"


# ---------------------------------------------------------------------------
# 项目加载与数据读取
# ---------------------------------------------------------------------------
def list_projects():
    """扫描 projects 目录下所有含 project.yaml 的项目。"""
    if not os.path.isdir(PROJECTS_DIR):
        return []
    projects = []
    for name in sorted(os.listdir(PROJECTS_DIR)):
        cfg_path = os.path.join(PROJECTS_DIR, name, "project.yaml")
        if os.path.isfile(cfg_path):
            projects.append(name)
    return projects


def load_project_config(project_id):
    """读取项目 project.yaml 配置。"""
    cfg_path = os.path.join(PROJECTS_DIR, project_id, "project.yaml")
    with open(cfg_path, "r", encoding="utf-8") as f:
        return yaml.safe_load(f)


def project_path(project_id, rel_path):
    """将 project.yaml 中的相对路径解析为绝对路径。"""
    return os.path.join(PROJECTS_DIR, project_id, rel_path)


@st.cache_data(show_spinner=False)
def load_csv(path):
    """安全加载 CSV/Excel 数据文件。"""
    if not os.path.isfile(path):
        return None
    try:
        if path.endswith((".xlsx", ".xls")):
            return pd.read_excel(path)
        return pd.read_csv(path, encoding="utf-8-sig")
    except Exception as e:
        st.warning(f"读取失败：{path}\n原因：{e}")
        return None


# ---------------------------------------------------------------------------
# 页面：使用指南
# ---------------------------------------------------------------------------
def page_user_guide(cfg, project_id):

    guide_path = os.path.join(
        BASE_DIR,
        "src",
        "user_guide.md",
    )

    if not os.path.isfile(guide_path):
        st.error("未找到使用指南文件。")
        return

    with open(
        guide_path,
        "r",
        encoding="utf-8",
    ) as guide_file:
        guide_text = guide_file.read()

    st.markdown(guide_text)

    st.download_button(
        "下载平台使用指南",
        data=guide_text.encode("utf-8-sig"),
        file_name="HerbInsight_平台使用指南.md",
        mime="text/markdown",
    )


# ---------------------------------------------------------------------------
# 页面：教学实训中心
# ---------------------------------------------------------------------------
def page_learning_center(cfg, project_id):
    st.header("教学实训中心")
    st.write(
        "从数据库检索开始，沿着标准科研流程边学边做。"
        "课程包括原理、资料模板、传统工具路线、参数与解释边界；"
        "现场计算、固定示意与尚未开放功能分别标注。"
    )

    st.caption("通用上传计算：网络边表分析、自备注释ORA；另有粘贴清单交集和多种固定教学演示。课程开放不等于算法上线，不使用课程数量推算算法完成比例。")
    st.info("从左侧「核心学习与实训」开始：先完成核心流程，再在「进阶证据」选择转录组、机器学习、免疫、MR或单细胞/扰动分支；最后进入「项目工具 → 资料模板与交付自查」。")

    stages = [
        ("1", "药材与成分收集", "记录数据库、版本、检索日期和筛选标准", "教学指南", "操作课程已开放"),
        ("2", "靶点标准化", "统一物种、UniProt ID 与基因符号并检查重复值", "数据治理", "操作课程已开放"),
        ("3", "疾病基因与交集", "整合 GeneCards、DisGeNET、OMIM 等多来源证据", "集合运算", "操作课程已开放"),
        ("4", "网络药理学", "构建药材—成分—靶点网络并计算拓扑指标", "真实计算", "已开放"),
        ("5", "GO/KEGG 富集", "理解超几何检验、背景集与 FDR，上传数据现场计算", "真实计算", "已开放"),
        ("6", "机器学习", "训练内筛选、多模型公平比较与内部保留测试集", "真实计算", "Logistic、Elastic Net、随机森林、SVM"),
        ("7", "孟德尔随机化", "外部LD筛选后，协调等位基因与IVW分析", "真实计算", "F筛选、IVW、Q和逐一剔除"),
        ("8", "免疫微环境证据", "上传表达与签名，比较相对富集分数", "真实计算", "ssGSEA、组间检验及BH"),
    ]

    for start in range(0, len(stages), 2):
        columns = st.columns(2)
        for column, (number, title, description, kind, status) in zip(
            columns, stages[start:start + 2]
        ):
            with column:
                with st.container(border=True):
                    st.subheader(f"{number}. {title}")
                    st.write(description)
                    tag1, tag2 = st.columns(2)
                    tag1.caption(f"类型：{kind}")
                    tag2.caption(f"状态：{status}")

    st.subheader("高级可选教学（不是必做流程）")
    advanced1, advanced2, advanced3 = st.columns(3)
    with advanced1:
        with st.container(border=True):
            st.subheader("9. 单细胞基础")
            st.write("质控、降维、聚类和细胞类型注释")
            st.caption("资料/QC/注释课程已开放；通用计算未开放")
    with advanced2:
        with st.container(border=True):
            st.subheader("10. 细胞通讯")
            st.write("基于单细胞数据和配体—受体库推断细胞互作")
            st.caption("课程已开放；不能由当前bulk直接完成；通用计算未开放")
    with advanced3:
        with st.container(border=True):
            st.subheader("11. 计算扰动")
            st.write("模拟基因干预后的网络或表达状态变化")
            st.caption("课程/网络删点示意已开放；真实表达扰动预测未开放")

    st.subheader("推荐学习路径")
    st.markdown(
        """
1. **先看五积散示范项目**：认识每一步的输入和输出。
2. **下载教学模板**：不要直接拿杂乱数据开始分析。
3. **运行小型示例**：理解参数变化怎样影响结果。
4. **上传自己的数据**：先通过字段和质量检查，再执行算法。
5. **完成结果解读**：回答“结果说明什么、不能说明什么”。
6. **导出实验记录**：保留数据来源、参数、方法和局限性。
        """
    )
    st.info(
        "现在可进入左侧“网络药理学教学实验”或“GO/KEGG 教学实验”，"
        "先运行内置示例理解输入、参数和输出，再替换成自己的数据。"
    )

# ---------------------------------------------------------------------------
# 页面：项目管理
# ---------------------------------------------------------------------------
def page_project_creation(cfg, project_id):
    from src.study_progress import render_project_creation
    render_project_creation()


def page_product_home(cfg, project_id):
    from src.product_home import render_product_home
    render_product_home(open_demo_workspace, lambda: navigate_to("learn_home"), open_wujisan_case)


def page_project_workspace(cfg, project_id):
    from src.project_workspace import render_project_workspace
    render_project_workspace(os.path.join(PROJECTS_DIR, project_id), cfg)


def page_project_management(cfg, project_id):
    st.header("项目与任务")
    st.caption("学习任务记录个人课程进度；数据项目保存可供工作台和案例页面读取的数据。")

    create_tab, upload_tab, status_tab = st.tabs(
        ["我的学习任务", "数据项目维护", "数据项目状态"]
    )
    with create_tab:
        st.write("新建学习任务后，按课程逐步完成资料收集与提交。已有任务可在侧栏切换。")
        if st.button("进入项目创建页", key="manager_open_creation"):
            st.query_params["view"] = "project_create"
            st.rerun()

    with upload_tab:
        st.subheader("数据项目维护")
        st.caption("这里维护服务器数据项目；个人学习资料请在对应课程的“本节任务”中提交。公开部署环境不用于保存私人研究数据。")

        available_projects = list_projects()

        if not available_projects:
            st.warning(
                "当前没有可用项目，请先创建项目。"
            )
        else:
            default_index = (
                available_projects.index(project_id)
                if project_id in available_projects
                else 0
            )

            upload_project_id = st.selectbox(
                "选择接收数据的项目",
                available_projects,
                index=default_index,
                key="upload_project_id",
            )

            file_labels = {
                key: (
                    f"{specification['description']} "
                    f"({key})"
                )
                for key, specification
                in project_manager
                .STANDARD_FILES.items()
                if key != "target_overlap"
            }

            selected_file_key = st.selectbox(
                "选择数据类型",
                list(file_labels.keys()),
                format_func=lambda key: file_labels[key],
            )

            specification = (
                project_manager.STANDARD_FILES[
                    selected_file_key
                ]
            )

            st.caption(
                "必需字段："
                + "、".join(
                    specification[
                        "required_columns"
                    ]
                )
            )

            template_table = pd.DataFrame(
                columns=specification[
                    "required_columns"
                ]
            )

            st.download_button(
                "下载当前数据类型的CSV模板",
                data=template_table.to_csv(
                    index=False,
                ).encode("utf-8-sig"),
                file_name=(
                    selected_file_key
                    + "_template.csv"
                ),
                mime="text/csv",
                key=(
                    "download_template_"
                    + selected_file_key
                ),
            )

            uploaded_file = st.file_uploader(
                "选择 CSV 或 Excel 文件",
                type=["csv", "xlsx", "xls"],
                key=(
                    "project_file_"
                    + selected_file_key
                ),
            )

            if st.button(
                "校验并保存数据",
                type="primary",
            ):
                if uploaded_file is None:
                    st.error("请先选择需要上传的文件。")
                else:
                    try:
                        result = (
                            project_manager
                            .save_uploaded_table(
                                project_id=(
                                    upload_project_id
                                ),
                                file_key=(
                                    selected_file_key
                                ),
                                file_name=(
                                    uploaded_file.name
                                ),
                                file_bytes=(
                                    uploaded_file.getvalue()
                                ),
                            )
                        )

                        load_csv.clear()

                        st.success(
                            f"{specification['description']}"
                            "上传成功。"
                        )

                        result1, result2 = st.columns(2)
                        result1.metric(
                            "记录数",
                            result["row_count"],
                        )
                        result2.metric(
                            "字段数",
                            result["column_count"],
                        )

                        st.write(
                            "**保存位置：** "
                            f"`{result['output_path']}`"
                        )

                    except (
                        ValueError,
                        KeyError,
                        FileNotFoundError,
                    ) as error:
                        st.error(str(error))

            st.divider()
            st.subheader("自动生成共同靶点")

            current_status = (
                project_manager
                .get_project_data_status(
                    upload_project_id
                )
            )

            uploaded_keys = set(
                current_status.loc[
                    current_status["status"]
                    == "已上传",
                    "file_key",
                ]
            )

            can_build_overlap = {
                "compound_targets",
                "disease_genes",
            }.issubset(uploaded_keys)

            if can_build_overlap:
                st.success(
                    "药材候选靶点与疾病基因已经齐全，"
                    "可以生成共同靶点。"
                )
            else:
                st.info(
                    "需要先上传“药材候选靶点”"
                    "和“疾病相关基因”。"
                )

            if st.button(
                "生成复方—疾病共同靶点",
                disabled=not can_build_overlap,
            ):
                try:
                    overlap = (
                        project_manager
                        .build_target_overlap(
                            upload_project_id
                        )
                    )

                    load_csv.clear()

                    st.success(
                        "共同靶点生成成功："
                        f"{len(overlap)} 个基因。"
                    )

                except (
                    ValueError,
                    KeyError,
                    FileNotFoundError,
                ) as error:
                    st.error(str(error))

    with status_tab:
        st.subheader("项目基础数据状态")

        available_projects = list_projects()

        if not available_projects:
            st.warning("当前没有可用项目。")
        else:
            default_index = (
                available_projects.index(project_id)
                if project_id in available_projects
                else 0
            )

            status_project_id = st.selectbox(
                "选择需要检查的项目",
                available_projects,
                index=default_index,
                key="status_project_id",
            )

            selected_config = (
                project_manager.load_project_config(
                    status_project_id
                )
            )

            status_data = (
                project_manager
                .get_project_data_status(
                    status_project_id
                )
            )

            completed_count = int(
                (
                    status_data["status"]
                    == "已上传"
                ).sum()
            )
            total_count = len(status_data)

            info1, info2, info3 = st.columns(3)

            info1.metric(
                "项目",
                selected_config.get(
                    "formula_name",
                    "-",
                ),
            )
            info2.metric(
                "疾病",
                selected_config.get(
                    "disease_name",
                    "-",
                ),
            )
            info3.metric(
                "基础数据",
                f"{completed_count}/{total_count}",
            )

            st.progress(
                completed_count / total_count
                if total_count
                else 0
            )

            st.dataframe(
                status_data,
                use_container_width=True,
                hide_index=True,
            )

            if completed_count == total_count:
                st.success(
                    "基础数据已经齐全，可以继续进行"
                    "PPI、富集和多组学分析。"
                )
            else:
                st.warning(
                    "项目仍有基础数据待上传。"
                )

    st.divider()
    st.info(
        "上传功能要求使用标准字段。平台会保留原始数据"
        "的分析含义，但会统一保存为UTF-8编码CSV文件。"
    )

# ---------------------------------------------------------------------------
# 页面：项目首页
# ---------------------------------------------------------------------------
def page_home(cfg, project_id):
    st.title(f"🌿 {cfg.get('project_name', '未命名项目')}")

    # 项目基本信息
    col1, col2, col3 = st.columns(3)
    col1.metric("复方", cfg.get("formula_name", "-"))
    col2.metric("疾病", cfg.get("disease_name", "-"))
    col3.metric("物种", cfg.get("species", "-"))

    st.caption(
        f"项目版本：{cfg.get('version', '-')} ｜ "
        f"创建日期：{cfg.get('created_date', '-')}"
    )

    st.divider()

    # 读取项目的标准化数据
    files = cfg.get("data_files", {})

    def get_data(file_key):
        rel_path = files.get(file_key)

        if not rel_path:
            return None

        return load_csv(
            project_path(project_id, rel_path)
        )

    herbs = get_data("herbs")
    targets = get_data("compound_targets")
    disease_genes = get_data("disease_genes")
    overlap = get_data("target_overlap")
    ppi_nodes = get_data("ppi_nodes")
    go_results = get_data("go_enrichment")
    kegg_results = get_data("kegg_enrichment")

    def row_count(data):
        return len(data) if data is not None else 0

    # 第一排统计卡片
    st.subheader("项目数据概览")

    card1, card2, card3, card4 = st.columns(4)

    card1.metric(
        "药材数量",
        f"{row_count(herbs):,}",
    )

    card2.metric(
        "候选靶点记录",
        f"{row_count(targets):,}",
    )

    card3.metric(
        "疾病基因",
        f"{row_count(disease_genes):,}",
    )

    card4.metric(
        "共同靶点",
        f"{row_count(overlap):,}",
    )

    # 第二排统计卡片
    card5, card6, card7, card8 = st.columns(4)

    card5.metric(
        "PPI节点",
        f"{row_count(ppi_nodes):,}",
    )

    card6.metric(
        "GO结果",
        f"{row_count(go_results):,}",
    )

    card7.metric(
        "KEGG结果",
        f"{row_count(kegg_results):,}",
    )

    modules = cfg.get("modules", [])

    implemented_modules = [
        module for module in modules
        if MODULE_FILES.get(module)
    ]
    completed_modules = sum(
        check_module_data(cfg, module)
        for module in implemented_modules
    )

    card8.metric(
        "已接入模块",
        f"{completed_modules}/{len(implemented_modules)}",
    )

    st.divider()

    # 模块状态
    st.subheader("模块完成状态")

    core_modules = [
        module for module in ["herbs", "target_overlap", "ppi", "enrichment"]
        if module in modules
    ]
    evidence_modules = [
        module for module in [
            "transcriptome", "machine_learning", "mendelian_randomization",
            "immune_infiltration", "integrated_evidence",
        ] if module in modules
    ]

    status1, status2 = st.columns(2)
    with status1:
        st.markdown("#### 核心网络药理学流程")
        for module in core_modules:
            status = "已接入数据" if check_module_data(cfg, module) else "待接入数据"
            st.write(f"- {MODULE_NAMES[module]}：{status}")
    with status2:
        st.markdown("#### 多组学证据扩展")
        for module in evidence_modules:
            status = "已接入数据" if check_module_data(cfg, module) else "待接入数据"
            st.write(f"- {MODULE_NAMES[module]}：{status}")

    st.info(
        "细胞通讯和计算扰动属于高级实验路线，目前尚未实现，"
        "因此不计入已接入模块。细胞通讯需要单细胞表达矩阵、"
        "细胞类型注释和配体—受体数据库。"
    )

    st.divider()
    st.info(cfg.get("disclaimer", DISCLAIMER))

# ---------------------------------------------------------------------------
# 模块名称与数据检查
# ---------------------------------------------------------------------------
MODULE_NAMES = {
    "herbs": "药材-候选靶点",
    "target_overlap": "疾病基因与交集靶点",
    "ppi": "网络药理学 / PPI",
    "enrichment": "GO/KEGG 富集分析",
    "transcriptome": "转录组验证",
    "machine_learning": "机器学习",
    "mendelian_randomization": "孟德尔随机化",
    "immune_infiltration": "免疫微环境证据",
    "integrated_evidence": "综合证据整合",
    "cell_communication": "细胞通讯",
    "virtual_knockout": "虚拟敲除",
}

# 每个模块依赖的数据文件（来自 project.yaml data_files）
MODULE_FILES = {
    "herbs": ["herbs", "compound_targets"],
    "target_overlap": ["disease_genes", "target_overlap"],
    "ppi": ["ppi_nodes", "ppi_edges"],
    "enrichment": ["go_enrichment", "kegg_enrichment"],
    "transcriptome": ["deg_results"],
    "machine_learning": ["ml_metrics"],
    "mendelian_randomization": ["mr_results"],
   "immune_infiltration": [
    "immune_scores",
    "immune_group_statistics",],
    "integrated_evidence": [
    "integrated_evidence_matrix",
    "integrated_evidence_summary",],
    "cell_communication": [],
    "virtual_knockout": [],
}


def check_module_data(cfg, mod):
    """检查模块要求的数据文件是否全部存在。"""
    files = cfg.get("data_files", {})
    required_files = MODULE_FILES.get(mod, [])

    if not required_files:
        return False

    for key in required_files:
        relative_path = files.get(key)

        if not relative_path:
            return False

        absolute_path = project_path(
            cfg["project_id"],
            relative_path,
        )

        if not os.path.isfile(absolute_path):
            return False

    return True


# ---------------------------------------------------------------------------
# 页面：药材-候选靶点
# ---------------------------------------------------------------------------
def page_herbs(cfg, project_id):
    st.header("药材与候选靶点")

    files = cfg.get("data_files", {})
    herbs = load_csv(
        project_path(project_id, files.get("herbs", ""))
    )
    ct = load_csv(
        project_path(project_id, files.get("compound_targets", ""))
    )

    # 药材清单
    if herbs is None:
        st.info("未找到 herbs.csv，请检查项目配置和数据文件。")
    else:
        with st.expander(f"查看药材清单（共 {len(herbs)} 味）"):
            st.dataframe(
                herbs,
                use_container_width=True,
                hide_index=True
            )

    # 候选靶点数据
    if ct is None:
        st.info("未找到候选靶点数据，请检查 compound_targets.csv。")
        return

    st.subheader("候选靶点筛选")

    filtered = ct.copy()

    # 第一排筛选条件
    col1, col2 = st.columns(2)

    if "herb_name" in ct.columns:
        herb_options = (
            ["全部药材"]
            + sorted(ct["herb_name"].dropna().astype(str).unique().tolist())
        )

        selected_herb = col1.selectbox(
            "选择药材",
            herb_options
        )

        if selected_herb != "全部药材":
            filtered = filtered[
                filtered["herb_name"].astype(str) == selected_herb
            ]
    else:
        selected_herb = "全部药材"
        col1.info("数据中没有 herb_name 字段")

    search = col2.text_input(
        "搜索基因或靶点",
        placeholder="例如：TNF、IL6、AKT1"
    )

    if search:
        search_mask = filtered.astype(str).apply(
            lambda row: row.str.contains(
                search,
                case=False,
                regex=False,
                na=False
            ).any(),
            axis=1
        )
        filtered = filtered[search_mask]

    # 第二排数值筛选条件
    col3, col4 = st.columns(2)

    fdr_limit = col3.number_input(
        "FDR 上限",
        min_value=0.0,
        max_value=1.0,
        value=1.0,
        step=0.01,
        help="数值越小，筛选越严格。1.0 表示暂不筛选。"
    )

    score_limit = col4.number_input(
        "证据分数下限",
        min_value=0.0,
        max_value=1.0,
        value=0.0,
        step=0.05,
        help="数值越大，筛选越严格。0.0 表示暂不筛选。"
    )

    if "fdr_bh" in filtered.columns and fdr_limit < 1.0:
        fdr_values = pd.to_numeric(
            filtered["fdr_bh"],
            errors="coerce"
        )
        filtered = filtered[fdr_values <= fdr_limit]

    if "evidence_score" in filtered.columns and score_limit > 0.0:
        score_values = pd.to_numeric(
            filtered["evidence_score"],
            errors="coerce"
        )
        filtered = filtered[score_values >= score_limit]

    # 结果统计
    total_records = len(ct)
    filtered_records = len(filtered)

    metric1, metric2, metric3 = st.columns(3)

    metric1.metric("全部候选记录", f"{total_records:,}")
    metric2.metric("当前筛选结果", f"{filtered_records:,}")

    if "gene_symbol" in filtered.columns:
        unique_genes = filtered["gene_symbol"].dropna().nunique()
    else:
        unique_genes = 0

    metric3.metric("当前候选基因", f"{unique_genes:,}")

    # 提供下载
    csv_data = filtered.to_csv(
        index=False
    ).encode("utf-8-sig")

    st.download_button(
        label="下载当前筛选结果（CSV）",
        data=csv_data,
        file_name="filtered_candidate_targets.csv",
        mime="text/csv"
    )

    st.caption(
        "说明：本页面展示的是数据库和计算分析得到的候选靶点，"
        "用于科研探索，不代表已经证实的治疗机制。"
    )

    st.dataframe(
        filtered,
        use_container_width=True,
        hide_index=True
    )

# ---------------------------------------------------------------------------
# 页面：疾病基因与交集靶点
# ---------------------------------------------------------------------------
def page_overlap(cfg, project_id):
    st.header("疾病基因与交集靶点")

    files = cfg.get("data_files", {})

    disease = load_csv(
        project_path(project_id, files.get("disease_genes", ""))
    )

    overlap = load_csv(
        project_path(project_id, files.get("target_overlap", ""))
    )

    # 页面总体统计
    disease_count = len(disease) if disease is not None else 0
    overlap_count = len(overlap) if overlap is not None else 0

    if disease is not None and "gene_symbol" in disease.columns:
        disease_gene_count = disease["gene_symbol"].dropna().nunique()
    else:
        disease_gene_count = 0

    if overlap is not None and "gene_symbol" in overlap.columns:
        overlap_gene_count = overlap["gene_symbol"].dropna().nunique()
    else:
        overlap_gene_count = 0

    metric1, metric2, metric3 = st.columns(3)

    metric1.metric("疾病基因记录", f"{disease_count:,}")
    metric2.metric("疾病基因数", f"{disease_gene_count:,}")
    metric3.metric("共同靶点数", f"{overlap_gene_count:,}")

    disease_tab, overlap_tab = st.tabs(
        ["疾病相关基因", "共同靶点"]
    )

    # 疾病相关基因
    with disease_tab:
        if disease is None:
            st.info("未找到 disease_genes.csv，请检查项目配置。")
        else:
            st.subheader("疾病相关基因筛选")

            disease_filtered = disease.copy()

            col1, col2 = st.columns(2)

            disease_search = col1.text_input(
                "搜索疾病基因",
                placeholder="例如：TNF、IL6、NOD2",
                key="disease_gene_search"
            )

            if "category" in disease.columns:
                category_options = (
                    ["全部类别"]
                    + sorted(
                        disease["category"]
                        .dropna()
                        .astype(str)
                        .unique()
                        .tolist()
                    )
                )

                selected_category = col2.selectbox(
                    "基因类别",
                    category_options,
                    key="disease_category"
                )

                if selected_category != "全部类别":
                    disease_filtered = disease_filtered[
                        disease_filtered["category"].astype(str)
                        == selected_category
                    ]
            else:
                selected_category = "全部类别"
                col2.info("数据中没有 category 字段")

            if disease_search:
                search_mask = disease_filtered.astype(str).apply(
                    lambda row: row.str.contains(
                        disease_search,
                        case=False,
                        regex=False,
                        na=False
                    ).any(),
                    axis=1
                )
                disease_filtered = disease_filtered[search_mask]

            if "relevance_score" in disease.columns:
                score_values = pd.to_numeric(
                    disease["relevance_score"],
                    errors="coerce"
                ).dropna()

                if not score_values.empty:
                    maximum_score = float(score_values.max())
                else:
                    maximum_score = 100.0

                minimum_score = st.number_input(
                    "相关性分数下限",
                    min_value=0.0,
                    max_value=max(maximum_score, 1.0),
                    value=0.0,
                    step=1.0,
                    help="数值越高，保留的疾病相关基因越少。",
                    key="minimum_relevance_score"
                )

                if minimum_score > 0:
                    numeric_scores = pd.to_numeric(
                        disease_filtered["relevance_score"],
                        errors="coerce"
                    )
                    disease_filtered = disease_filtered[
                        numeric_scores >= minimum_score
                    ]

            st.metric(
                "当前疾病基因结果",
                f"{len(disease_filtered):,}"
            )

            disease_csv = disease_filtered.to_csv(
                index=False
            ).encode("utf-8-sig")

            st.download_button(
                "下载当前疾病基因结果（CSV）",
                data=disease_csv,
                file_name="filtered_disease_genes.csv",
                mime="text/csv",
                key="download_disease_genes"
            )

            st.dataframe(
                disease_filtered,
                use_container_width=True,
                hide_index=True
            )

    # 共同靶点
    with overlap_tab:
        if overlap is None:
            st.info(
                "未找到 target_overlap.csv，请检查交集分析结果。"
            )
        else:
            st.subheader("药物候选靶点与疾病基因的交集")

            overlap_filtered = overlap.copy()

            col3, col4 = st.columns(2)

            overlap_search = col3.text_input(
                "搜索共同靶点",
                placeholder="例如：TNF、IL6、AKT1",
                key="overlap_gene_search"
            )

            if "evidence_type" in overlap.columns:
                evidence_options = (
                    ["全部证据类型"]
                    + sorted(
                        overlap["evidence_type"]
                        .dropna()
                        .astype(str)
                        .unique()
                        .tolist()
                    )
                )

                selected_evidence = col4.selectbox(
                    "证据类型",
                    evidence_options,
                    key="overlap_evidence_type"
                )

                if selected_evidence != "全部证据类型":
                    overlap_filtered = overlap_filtered[
                        overlap_filtered["evidence_type"].astype(str)
                        == selected_evidence
                    ]
            else:
                selected_evidence = "全部证据类型"
                col4.info("数据中没有 evidence_type 字段")

            if overlap_search:
                overlap_mask = overlap_filtered.astype(str).apply(
                    lambda row: row.str.contains(
                        overlap_search,
                        case=False,
                        regex=False,
                        na=False
                    ).any(),
                    axis=1
                )
                overlap_filtered = overlap_filtered[overlap_mask]

            if "gene_symbol" in overlap_filtered.columns:
                current_overlap_genes = (
                    overlap_filtered["gene_symbol"]
                    .dropna()
                    .nunique()
                )
            else:
                current_overlap_genes = len(overlap_filtered)

            result1, result2 = st.columns(2)

            result1.metric(
                "当前共同靶点记录",
                f"{len(overlap_filtered):,}"
            )

            result2.metric(
                "当前共同基因数",
                f"{current_overlap_genes:,}"
            )

            overlap_csv = overlap_filtered.to_csv(
                index=False
            ).encode("utf-8-sig")

            st.download_button(
                "下载当前共同靶点（CSV）",
                data=overlap_csv,
                file_name="filtered_overlap_targets.csv",
                mime="text/csv",
                key="download_overlap_targets"
            )

            st.caption(
                "说明：共同靶点表示药物候选靶点与疾病相关基因的集合交集，"
                "用于筛选后续研究对象，并不等同于已经验证的治疗机制。"
            )

            st.dataframe(
                overlap_filtered,
                use_container_width=True,
                hide_index=True
            )

# ---------------------------------------------------------------------------
# 页面：网络药理学 / PPI
# ---------------------------------------------------------------------------
def page_ppi_results(cfg, project_id):
    import networkx as nx
    import plotly.graph_objects as go

    st.header("网络药理学 / PPI 网络")

    files = cfg.get("data_files", {})

    nodes = load_csv(
        project_path(project_id, files.get("ppi_nodes", ""))
    )

    edges = load_csv(
        project_path(project_id, files.get("ppi_edges", ""))
    )

    if nodes is None or edges is None:
        st.info(
            "尚未找到完整的 PPI 节点和互作边数据。"
            "请检查 ppi_nodes.csv 与 ppi_edges.csv。"
        )
        return

    required_node_columns = {
        "gene_symbol",
        "degree",
        "betweenness",
        "closeness",
    }

    required_edge_columns = {
        "source",
        "target",
        "score",
    }

    missing_node_columns = (
        required_node_columns - set(nodes.columns)
    )

    missing_edge_columns = (
        required_edge_columns - set(edges.columns)
    )

    if missing_node_columns or missing_edge_columns:
        st.error(
            "PPI 数据字段不完整。"
            f"节点缺失字段：{sorted(missing_node_columns)}；"
            f"边缺失字段：{sorted(missing_edge_columns)}"
        )
        return

    # 数据类型标准化
    nodes = nodes.copy()
    edges = edges.copy()

    nodes["gene_symbol"] = (
        nodes["gene_symbol"].astype(str)
    )

    nodes["degree"] = pd.to_numeric(
        nodes["degree"],
        errors="coerce"
    ).fillna(0)

    nodes["betweenness"] = pd.to_numeric(
        nodes["betweenness"],
        errors="coerce"
    ).fillna(0)

    nodes["closeness"] = pd.to_numeric(
        nodes["closeness"],
        errors="coerce"
    ).fillna(0)

    edges["source"] = edges["source"].astype(str)
    edges["target"] = edges["target"].astype(str)

    edges["score"] = pd.to_numeric(
        edges["score"],
        errors="coerce"
    )

    edges = edges.dropna(
        subset=["source", "target", "score"]
    )

    nodes = nodes.sort_values(
        ["degree", "betweenness"],
        ascending=[False, False]
    ).reset_index(drop=True)

    # 总体统计
    maximum_degree = int(nodes["degree"].max())

    metric1, metric2, metric3, metric4 = st.columns(4)

    metric1.metric(
        "PPI节点数",
        f"{len(nodes):,}"
    )

    metric2.metric(
        "高置信度互作边",
        f"{len(edges):,}"
    )

    metric3.metric(
        "最高 Degree",
        f"{maximum_degree:,}"
    )

    metric4.metric(
        "STRING阈值",
        "≥ 0.700"
    )

    st.caption(
        "数据来源：STRING；物种：Homo sapiens；"
        "网络类型：full STRING network；"
        "最低互作分数：0.700。"
    )

    st.subheader("核心靶点与子网络设置")

    control1, control2 = st.columns(2)

    top_n = control1.slider(
        "显示 Degree 排名前 N 个靶点",
        min_value=10,
        max_value=len(nodes),
        value=min(30, len(nodes)),
        step=5,
        help="推荐先查看前30个核心靶点；增大数量会使网络更加密集。"
    )

    minimum_score = control2.slider(
        "当前网络最低互作分数",
        min_value=0.70,
        max_value=1.00,
        value=0.70,
        step=0.01,
        help="数值越高，保留的互作关系越可靠，但边数会减少。"
    )

    # 选取 Degree 排名前 N 个节点
    top_nodes = nodes.head(top_n).copy()

    selected_genes = set(
        top_nodes["gene_symbol"].tolist()
    )

    sub_edges = edges[
        edges["source"].isin(selected_genes)
        & edges["target"].isin(selected_genes)
        & (edges["score"] >= minimum_score)
    ].copy()

    result1, result2 = st.columns(2)

    result1.metric(
        "当前核心靶点数",
        f"{len(top_nodes):,}"
    )

    result2.metric(
        "当前子网络互作边",
        f"{len(sub_edges):,}"
    )

    ranking_tab, network_tab = st.tabs(
        ["核心靶点排名", "核心子网络"]
    )

    # 核心靶点排名
    with ranking_tab:
        st.subheader("核心靶点拓扑指标")

        gene_search = st.text_input(
            "搜索核心靶点",
            placeholder="例如：AKT1、TNF、IL6",
            key="ppi_gene_search"
        )

        ranking_display = top_nodes.copy()

        if gene_search:
            ranking_mask = (
                ranking_display.astype(str).apply(
                    lambda row: row.str.contains(
                        gene_search,
                        case=False,
                        regex=False,
                        na=False
                    ).any(),
                    axis=1
                )
            )

            ranking_display = ranking_display[
                ranking_mask
            ]

        ranking_columns = [
            column
            for column in [
                "gene_symbol",
                "degree",
                "degree_centrality",
                "betweenness",
                "closeness",
                "connected",
            ]
            if column in ranking_display.columns
        ]

        st.dataframe(
            ranking_display[ranking_columns],
            use_container_width=True,
            hide_index=True
        )

        nodes_csv = top_nodes.to_csv(
            index=False
        ).encode("utf-8-sig")

        st.download_button(
            "下载当前核心靶点排名（CSV）",
            data=nodes_csv,
            file_name=f"ppi_top_{top_n}_nodes.csv",
            mime="text/csv",
            key="download_ppi_nodes"
        )

    # 核心子网络
    with network_tab:
        st.subheader(
            f"Degree Top {top_n} 核心靶点子网络"
        )

        graph = nx.Graph()

        graph.add_nodes_from(
            top_nodes["gene_symbol"].tolist()
        )

        for _, row in sub_edges.iterrows():
            graph.add_edge(
                str(row["source"]),
                str(row["target"]),
                weight=float(row["score"])
            )

        position = nx.spring_layout(
            graph,
            seed=42,
            k=0.8,
            iterations=100,
            weight="weight"
        )

        # 绸制网络边
        edge_x = []
        edge_y = []

        for source, target in graph.edges():
            x0, y0 = position[source]
            x1, y1 = position[target]

            edge_x.extend([x0, x1, None])
            edge_y.extend([y0, y1, None])

        edge_trace = go.Scatter(
            x=edge_x,
            y=edge_y,
            mode="lines",
            line=dict(
                width=1.0,
                color="rgba(125, 211, 166, 0.45)"
            ),
            hoverinfo="none",
            name="互作边"
        )

        node_lookup = top_nodes.set_index(
            "gene_symbol"
        )

        node_x = []
        node_y = []
        node_text = []
        node_hover = []
        node_degree = []
        node_size = []

        highest_sub_degree = max(
            [
                graph.degree(node)
                for node in graph.nodes()
            ]
            or [1]
        )

        for node in graph.nodes():
            x, y = position[node]
            node_x.append(x)
            node_y.append(y)
            node_text.append(node)

            degree_value = int(
                node_lookup.loc[node, "degree"]
            )

            betweenness_value = float(
                node_lookup.loc[node, "betweenness"]
            )

            closeness_value = float(
                node_lookup.loc[node, "closeness"]
            )

            node_degree.append(degree_value)

            current_sub_degree = graph.degree(node)

            node_size.append(
                14
                + 24
                * current_sub_degree
                / max(highest_sub_degree, 1)
            )

            node_hover.append(
                f"<b>{node}</b><br>"
                f"全网 Degree：{degree_value}<br>"
                f"当前子网 Degree：{current_sub_degree}<br>"
                f"Betweenness：{betweenness_value:.4f}<br>"
                f"Closeness：{closeness_value:.4f}"
            )

        node_trace = go.Scatter(
            x=node_x,
            y=node_y,
            mode="markers+text",
            text=node_text,
            textposition="top center",
            textfont=dict(size=11),
            hovertext=node_hover,
            hoverinfo="text",
            marker=dict(
                size=node_size,
                color=node_degree,
                colorscale="Viridis",
                showscale=True,
                colorbar=dict(
                    title="Degree",
                    thickness=14
                ),
                line=dict(
                    width=1,
                    color="#d9f5e5"
                )
            ),
            name="核心靶点"
        )

        figure = go.Figure(
            data=[edge_trace, node_trace]
        )

        figure.update_layout(
            showlegend=False,
            hovermode="closest",
            height=720,
            margin=dict(
                l=20,
                r=20,
                t=60,
                b=20
            ),
            title=(
                f"PPI核心子网络：Top {top_n} 节点，"
                f"{len(sub_edges)} 条互作边"
            ),
            xaxis=dict(
                showgrid=False,
                zeroline=False,
                showticklabels=False
            ),
            yaxis=dict(
                showgrid=False,
                zeroline=False,
                showticklabels=False
            )
        )

        st.plotly_chart(
            figure,
            use_container_width=True
        )

        edges_csv = sub_edges.to_csv(
            index=False
        ).encode("utf-8-sig")

        st.download_button(
            "下载当前核心子网络互作边（CSV）",
            data=edges_csv,
            file_name=f"ppi_top_{top_n}_edges.csv",
            mime="text/csv",
            key="download_ppi_edges"
        )

    st.caption(
        "科学解释：Degree、Betweenness 和 Closeness "
        "用于描述网络拓扑位置，可辅助确定候选核心靶点；"
        "拓扑中心性不代表治疗因果关系，仍需实验或多组学证据验证。"
    )


def _plot_uploaded_network(graph, nodes, title):
    import networkx as nx
    import plotly.graph_objects as go

    position = nx.spring_layout(graph, seed=42, weight="weight", iterations=100)
    edge_x, edge_y = [], []
    for source, target in graph.edges():
        x0, y0 = position[source]
        x1, y1 = position[target]
        edge_x.extend([x0, x1, None])
        edge_y.extend([y0, y1, None])
    edge_trace = go.Scatter(
        x=edge_x, y=edge_y, mode="lines",
        line={"width": 1, "color": "rgba(91, 125, 111, 0.35)"},
        hoverinfo="none", name="关系", showlegend=False,
    )

    lookup = nodes.set_index("node_id")
    node_types = sorted(nodes["node_type"].astype(str).unique())
    colors = ["#2E7D5B", "#D18B35", "#3E75A6", "#A65353", "#7A65A8", "#64746D"]
    traces = [edge_trace]
    for index, node_type in enumerate(node_types):
        type_nodes = [
            node for node in graph.nodes()
            if str(lookup.loc[node, "node_type"]) == node_type
        ]
        hover = [
            f"<b>{node}</b><br>类型：{node_type}<br>"
            f"Degree：{int(lookup.loc[node, 'degree'])}<br>"
            f"Betweenness：{lookup.loc[node, 'betweenness']:.4f}<br>"
            f"Closeness：{lookup.loc[node, 'closeness']:.4f}"
            for node in type_nodes
        ]
        traces.append(
            go.Scatter(
                x=[position[node][0] for node in type_nodes],
                y=[position[node][1] for node in type_nodes],
                mode="markers+text", text=type_nodes, textposition="top center",
                hovertext=hover, hoverinfo="text", name=node_type,
                marker={
                    "size": [14 + 4 * int(lookup.loc[node, "degree"]) for node in type_nodes],
                    "color": colors[index % len(colors)],
                    "line": {"width": 1, "color": "white"},
                },
            )
        )
    figure = go.Figure(traces)
    figure.update_layout(
        title=title, height=680, hovermode="closest",
        margin={"l": 10, "r": 10, "t": 55, "b": 10},
        xaxis={"visible": False}, yaxis={"visible": False},
        legend={"title": "节点类型"},
    )
    st.plotly_chart(figure, use_container_width=True)


def _network_method_text(diagnostics, minimum_weight, directed):
    network_type = "有向网络" if directed else "无向网络"
    return f"""# 网络药理学在线实验记录

## 方法
根据上传的边表构建{network_type}。删除自环，以节点对为单位合并重复关系并保留最高权重，随后计算 Degree、Degree Centrality、Betweenness Centrality、Closeness Centrality、连通分量和网络密度。

## 数据与参数
- 原始有效边记录：{diagnostics['uploaded_edge_count']}
- 最低权重阈值：{minimum_weight}
- 最终节点：{diagnostics['node_count']}
- 最终边：{diagnostics['edge_count']}
- 重复边：{diagnostics['duplicate_count']}
- 自环：{diagnostics['self_loop_count']}
- 连通分量：{diagnostics['component_count']}
- 网络密度：{diagnostics['density']:.6f}

## 解释边界
中心性按无权最短路径计算；权重仅用于阈值筛选和布局。有向图Closeness使用指向节点的距离。重复节点对合并为简单图，不保留多重边；仅有边表无法恢复孤立节点。
中心性衡量节点在当前网络结构中的位置，受数据来源、纳入规则和网络完整性影响。中心节点是优先研究候选，不等于关键治疗靶点，也不能单独证明因果关系。
"""


def page_network_lab(cfg, project_id, embedded=False):
    st.subheader("上传边表并计算" if embedded else "网络药理学在线实验室")
    st.write(
        "上传药材—成分—靶点或蛋白互作边表，平台将现场清洗关系、"
        "构建网络并计算拓扑指标。"
    )
    learn_tab, prepare_tab, calculate_tab, interpret_tab = (
        (st.expander("查看指标解释"), st.expander("输入格式与示例下载"), st.container(), st.expander("怎样解读结果"))
        if embedded else st.tabs(
        ["① 学原理", "② 备数据", "③ 做分析", "④ 读结果"]
    ))
    with learn_tab:
        st.markdown(
            """
### 网络药理学在做什么？

网络把研究对象表示为**节点**，把它们之间的关系表示为**边**。
药材—成分—靶点网络用于整理多层关联，PPI 网络用于观察蛋白之间的互作结构。

- **Degree**：直接相连的边数，反映局部连接数量；
- **Degree Centrality**：Degree 按网络规模标准化后的结果；
- **Betweenness**：节点位于其他节点最短路径上的程度；
- **Closeness**：节点到网络中其他节点的平均距离是否较短；
- **连通分量**：彼此可到达的一组节点，可用于发现孤立子网络。
            """
        )
        st.warning(
            "中心性高只表示节点在当前网络中位置突出。数据库收录偏倚、"
            "网络阈值和研究较多的基因都会影响排名。"
        )

    example_content = example_edge_file()
    with prepare_tab:
        st.info("边表准备见本节前面的STRING下载与Cytoscape操作说明。此处计算使用你已获取的关系数据。")
        st.warning("内置网络为教学示意关系与虚构权重，未逐条核实数据库证据，不可作为真实研究结果。")
        st.markdown(
            """
### 边表格式

每行表示一条关系。必要字段为 `source` 和 `target`；推荐增加：

- `source_type`、`target_type`：如复方、药材、成分、靶点；
- `relation`：如包含、含有、作用或 PPI；
- `weight`：关系置信度，未提供时默认为 1。

同名节点会被视为同一个实体，因此应在上传前统一药材名、成分 ID 和基因符号。
            """
        )
        st.download_button(
            "下载网络边表示例（CSV）", example_content,
            file_name="network_edges_example.csv", mime="text/csv",
            use_container_width=True,
        )
        st.dataframe(
            pd.read_csv(BytesIO(example_content)),
            use_container_width=True, hide_index=True,
        )

    with calculate_tab:
        use_example = st.toggle(
            "使用内置教学示例", value=True, key="network_use_example",
            help="首次学习建议使用内置的复方—药材—成分—靶点小型网络。",
        )
        uploaded_edges = None
        edge_data = None
        mapping_record = "内置教学示意数据；非数据库证据"
        if not use_example:
            uploaded_edges = st.file_uploader(
                "上传网络边表", type=["csv", "tsv", "txt", "xlsx", "xls"],
                key="network_edge_upload",
            )
            if uploaded_edges is not None:
                try:
                    raw_edges = read_edge_table(uploaded_edges)
                    if not raw_edges.columns.is_unique:
                        raise ValueError("列名不能重复。")
                    columns = list(raw_edges.columns)
                    if len(columns) < 2:
                        raise ValueError("至少需要两列端点；TSV请使用.tsv或.txt后缀。")
                    def guess(candidates, fallback):
                        return next((i for i, c in enumerate(columns) if str(c).lower().strip() in candidates), fallback)
                    st.dataframe(raw_edges.head(10), hide_index=True, use_container_width=True)
                    source_col = st.selectbox("起点列（source）", columns, index=guess(["source", "from", "#node1", "node1", "preferredname_a", "stringid_a"], 0))
                    target_col = st.selectbox("终点列（target）", columns, index=guess(["target", "to", "node2", "preferredname_b", "stringid_b"], 1))
                    weight_options = [None] + columns
                    weight_col = st.selectbox("权重/置信度列", weight_options, index=guess(["weight", "score", "combined_score", "combinedscore", "confidence"], -1) + 1, format_func=lambda x: "不提供（默认1）" if x is None else str(x))
                    scale = st.selectbox("分数尺度（请核对导出说明）", ["保持原始权重", "STRING 0–1000 转为0–1"])
                    edge_data = map_edge_columns(raw_edges, source_col, target_col, weight_col, "string_1000" if scale.startswith("STRING") else "original")
                    mapping_record = f"文件：{uploaded_edges.name}；起点：{source_col}；终点：{target_col}；权重：{weight_col}；尺度：{scale}"
                    st.caption("下方阈值使用转换后的单位。STRING转换后0.7对应原始700；节点仍使用所选列的原始ID。在线上限500节点/5000条去重边。")
                except Exception as error:
                    st.error(f"边表准备未完成：{error}")
        option1, option2 = st.columns(2)
        minimum_weight = option1.number_input(
            "最低关系权重", min_value=0.0, value=0.0, step=0.05,
            help="低于该数值的边不会进入网络。无权重数据可保持为0。",
        )
        directed = option2.toggle(
            "构建有向网络", value=False,
            help="药材→成分等方向有明确含义时可开启；PPI通常使用无向网络。",
        )

        if st.button("开始构建并分析网络", type="primary", use_container_width=True):
            try:
                if use_example:
                    edge_data = pd.read_csv(BytesIO(example_content))
                elif edge_data is None:
                    raise ValueError("请上传网络边表并正确选择列与尺度。")
                nodes, edges, graph, diagnostics = analyze_network(
                    edge_data, float(minimum_weight), bool(directed)
                )
                st.session_state["network_nodes"] = nodes
                st.session_state["network_edges"] = edges
                st.session_state["network_graph"] = graph
                st.session_state["network_diagnostics"] = diagnostics
                st.session_state["network_parameters"] = {
                    "minimum_weight": float(minimum_weight), "directed": bool(directed),
                    "mapping_record": mapping_record,
                }
            except Exception as error:
                for result_key in ["network_nodes", "network_edges", "network_graph", "network_diagnostics", "network_parameters"]:
                    st.session_state.pop(result_key, None)
                st.error(f"网络分析未完成：{error}")

        nodes = st.session_state.get("network_nodes")
        edges = st.session_state.get("network_edges")
        graph = st.session_state.get("network_graph")
        diagnostics = st.session_state.get("network_diagnostics")
        parameters = st.session_state.get("network_parameters", {})
        if nodes is not None and diagnostics is not None and graph is not None:
            st.success("网络构建完成。以下指标由本次边表现场计算。")
            st.caption("已计算输入：" + parameters.get("mapping_record", "旧会话数据") + "。修改输入或参数后请重新运行；有向图的线段不显示箭头，以边表方向为准。")
            metrics = st.columns(4)
            metrics[0].metric("节点", diagnostics["node_count"])
            metrics[1].metric("关系边", diagnostics["edge_count"])
            metrics[2].metric("连通分量", diagnostics["component_count"])
            metrics[3].metric("网络密度", f"{diagnostics['density']:.3f}")

            if diagnostics["duplicate_count"] or diagnostics["self_loop_count"]:
                st.info(
                    f"质量检查：合并重复边 {diagnostics['duplicate_count']} 条，"
                    f"删除自环 {diagnostics['self_loop_count']} 条。"
                )
            result_tab, graph_tab = st.tabs(["拓扑指标排名", "交互网络"])
            with result_tab:
                st.dataframe(nodes, use_container_width=True, hide_index=True)
            with graph_tab:
                _plot_uploaded_network(graph, nodes, "本次上传数据构建的网络")

            download1, download2 = st.columns(2)
            download1.download_button(
                "下载节点拓扑指标（CSV）",
                nodes.to_csv(index=False).encode("utf-8-sig"),
                "network_node_metrics.csv", "text/csv", use_container_width=True,
            )
            download2.download_button(
                "下载清洗后的边表（CSV）",
                edges.to_csv(index=False).encode("utf-8-sig"),
                "network_clean_edges.csv", "text/csv", use_container_width=True,
            )
            method_text = _network_method_text(
                diagnostics,
                parameters.get("minimum_weight", minimum_weight),
                parameters.get("directed", directed),
            )
            method_text += "\n## 输入列与尺度\n" + parameters.get("mapping_record", "未记录")
            st.download_button(
                "下载实验方法与质量记录（Markdown）",
                method_text.encode("utf-8-sig"), "network_experiment_record.md",
                "text/markdown", use_container_width=True,
            )

    with interpret_tab:
        st.markdown(
            """
### 结果解读顺序

1. 先看清洗提示，确认重复边、自环和权重阈值是否符合预期；
2. 检查连通分量，避免把多个互不相连的子网络混为一个机制；
3. 比较 Degree 与 Betweenness，识别“连接多”和“桥梁型”节点；
4. 回查核心节点由哪些原始关系支持，而不是只报告排名；
5. 使用不同阈值重复分析，观察核心节点是否稳定；
6. 将网络结果与转录组、MR或实验结果交叉验证。

### 课堂任务

- 把最低权重逐步提高，网络密度和核心节点如何变化？
- 同一张边表分别构建有向和无向网络，Degree 的含义有何不同？
- 为什么文献研究较多的蛋白可能更容易成为网络中心？
            """
        )


def page_ppi(cfg, project_id):
    st.header("网络药理学教学实验")
    page_network_lab(cfg, project_id)

# ---------------------------------------------------------------------------
# 页面：GO/KEGG 富集分析
# ---------------------------------------------------------------------------
def page_enrichment_results(cfg, project_id):
    st.header("GO / KEGG 功能富集分析")

    files = cfg.get("data_files", {})

    go_df = load_csv(
        project_path(
            project_id,
            files.get("go_enrichment", "")
        )
    )

    kegg_df = load_csv(
        project_path(
            project_id,
            files.get("kegg_enrichment", "")
        )
    )

    # 自动识别校正后 P 值字段
    def find_adjusted_p_column(data):
        if data is None:
            return None

        candidates = [
            "p_adjust",
            "p.adjust",
            "p_adj",
            "adjusted_p_value",
            "adjusted_p",
            "fdr_bh",
            "qvalue",
            "q_value",
            "FDR",
        ]

        for column in candidates:
            if column in data.columns:
                return column

        lower_columns = {
            str(column).lower(): column
            for column in data.columns
        }

        for candidate in candidates:
            if candidate.lower() in lower_columns:
                return lower_columns[candidate.lower()]

        return None

    go_count = len(go_df) if go_df is not None else 0
    kegg_count = len(kegg_df) if kegg_df is not None else 0

    metric1, metric2, metric3 = st.columns(3)

    metric1.metric(
        "GO富集结果",
        f"{go_count:,}"
    )

    metric2.metric(
        "KEGG富集结果",
        f"{kegg_count:,}"
    )

    metric3.metric(
        "输入共同靶点",
        "191"
    )

    st.caption(
        "富集分析用于识别共同靶点可能参与的生物学功能和信号通路。"
        "结果反映统计关联，不等同于已经验证的作用机制。"
    )

    tab_go, tab_kegg = st.tabs(
        ["GO富集", "KEGG通路"]
    )

    # -------------------------
    # GO富集
    # -------------------------
    with tab_go:
        if go_df is None:
            st.info("未找到 go_enrichment.csv。")
        else:
            st.subheader("GO富集结果筛选")

            go_filtered = go_df.copy()

            go_search_col, ontology_filter_col, go_fdr_col = (
                st.columns(3)
            )

            go_search = go_search_col.text_input(
                "搜索GO条目",
                placeholder="例如：inflammatory、oxidative",
                key="go_term_search"
            )

            ontology_column = next(
                (
                    column
                    for column in go_df.columns
                    if str(column).lower() == "ontology"
                ),
                None
            )

            if ontology_column is not None:
                ontology_options = (
                    ["全部类型"]
                    + sorted(
                        go_df[ontology_column]
                        .dropna()
                        .astype(str)
                        .unique()
                        .tolist()
                    )
                )

                selected_ontology = (
                    ontology_filter_col.selectbox(
                        "GO本体类型",
                        ontology_options,
                        key="go_ontology_filter"
                    )
                )

                if selected_ontology != "全部类型":
                    go_filtered = go_filtered[
                        go_filtered[
                            ontology_column
                        ].astype(str)
                        == selected_ontology
                    ]
            else:
                selected_ontology = "全部类型"
                ontology_filter_col.info(
                    "未找到 ONTOLOGY 字段"
                )

            go_fdr_limit = go_fdr_col.number_input(
                "校正后P值上限",
                min_value=0.0,
                max_value=1.0,
                value=0.05,
                step=0.001,
                format="%.4f",
                key="go_fdr_limit"
            )

            if go_search:
                go_search_columns = [
                    column
                    for column in [
                        "term_id",
                        "term_name",
                    ]
                    if column in go_filtered.columns
                ]

                if go_search_columns:
                    go_search_mask = (
                        go_filtered[
                            go_search_columns
                        ]
                        .astype(str)
                        .apply(
                            lambda row: row.str.contains(
                                go_search,
                                case=False,
                                regex=False,
                                na=False
                            ).any(),
                            axis=1
                        )
                    )

                    go_filtered = go_filtered[
                        go_search_mask
                    ]

            go_adjusted_p = find_adjusted_p_column(
                go_filtered
            )

            if go_adjusted_p is not None:
                go_p_values = pd.to_numeric(
                    go_filtered[go_adjusted_p],
                    errors="coerce"
                )

                go_filtered = go_filtered[
                    go_p_values <= go_fdr_limit
                ]

                go_filtered = go_filtered.sort_values(
                    go_adjusted_p,
                    ascending=True
                )
            else:
                st.warning(
                    "未识别到校正后P值字段，"
                    "本次不执行显著性筛选。"
                )

            go_result1, go_result2 = st.columns(2)

            go_result1.metric(
                "全部GO条目",
                f"{len(go_df):,}"
            )

            go_result2.metric(
                "当前筛选结果",
                f"{len(go_filtered):,}"
            )

            go_csv = go_filtered.to_csv(
                index=False
            ).encode("utf-8-sig")

            st.download_button(
                "下载当前GO筛选结果（CSV）",
                data=go_csv,
                file_name="filtered_go_enrichment.csv",
                mime="text/csv",
                key="download_go_results"
            )

            st.dataframe(
                go_filtered,
                use_container_width=True,
                hide_index=True
            )

            if not go_filtered.empty:
                _plot_bubble(
                    go_filtered,
                    "GO富集气泡图"
                )
            else:
                st.warning(
                    "当前条件下没有GO富集结果，"
                    "请放宽筛选条件。"
                )

    # -------------------------
    # KEGG富集
    # -------------------------
    with tab_kegg:
        if kegg_df is None:
            st.info("未找到 kegg_enrichment.csv。")
        else:
            st.subheader("KEGG通路结果筛选")

            kegg_filtered = kegg_df.copy()

            kegg_search_col, kegg_fdr_col = (
                st.columns(2)
            )

            kegg_search = kegg_search_col.text_input(
                "搜索KEGG通路",
                placeholder="例如：TNF、PI3K、apoptosis",
                key="kegg_term_search"
            )

            kegg_fdr_limit = (
                kegg_fdr_col.number_input(
                    "校正后P值上限",
                    min_value=0.0,
                    max_value=1.0,
                    value=0.05,
                    step=0.001,
                    format="%.4f",
                    key="kegg_fdr_limit"
                )
            )

            if kegg_search:
                kegg_search_columns = [
                    column
                    for column in [
                        "term_id",
                        "term_name",
                    ]
                    if column in kegg_filtered.columns
                ]

                if kegg_search_columns:
                    kegg_search_mask = (
                        kegg_filtered[
                            kegg_search_columns
                        ]
                        .astype(str)
                        .apply(
                            lambda row: row.str.contains(
                                kegg_search,
                                case=False,
                                regex=False,
                                na=False
                            ).any(),
                            axis=1
                        )
                    )

                    kegg_filtered = kegg_filtered[
                        kegg_search_mask
                    ]

            kegg_adjusted_p = find_adjusted_p_column(
                kegg_filtered
            )

            if kegg_adjusted_p is not None:
                kegg_p_values = pd.to_numeric(
                    kegg_filtered[
                        kegg_adjusted_p
                    ],
                    errors="coerce"
                )

                kegg_filtered = kegg_filtered[
                    kegg_p_values <= kegg_fdr_limit
                ]

                kegg_filtered = (
                    kegg_filtered.sort_values(
                        kegg_adjusted_p,
                        ascending=True
                    )
                )
            else:
                st.warning(
                    "未识别到校正后P值字段，"
                    "本次不执行显著性筛选。"
                )

            kegg_result1, kegg_result2 = (
                st.columns(2)
            )

            kegg_result1.metric(
                "全部KEGG通路",
                f"{len(kegg_df):,}"
            )

            kegg_result2.metric(
                "当前筛选结果",
                f"{len(kegg_filtered):,}"
            )

            kegg_csv = kegg_filtered.to_csv(
                index=False
            ).encode("utf-8-sig")

            st.download_button(
                "下载当前KEGG筛选结果（CSV）",
                data=kegg_csv,
                file_name="filtered_kegg_enrichment.csv",
                mime="text/csv",
                key="download_kegg_results"
            )

            st.dataframe(
                kegg_filtered,
                use_container_width=True,
                hide_index=True
            )

            if not kegg_filtered.empty:
                _plot_bubble(
                    kegg_filtered,
                    "KEGG富集气泡图"
                )
            else:
                st.warning(
                    "当前条件下没有KEGG通路结果，"
                    "请放宽筛选条件。"
                )
   

def _enrichment_method_text(diagnostics, minimum_overlap, fdr_limit):
    return f"""# GO/KEGG 过度富集分析实验记录

## 方法
以用户上传的目标基因作为样本，以背景基因作为抽样总体，使用单侧超几何检验评估每个功能条目是否出现过度富集，并采用 Benjamini-Hochberg 方法控制多重检验的错误发现率。
先对所有实际检验条目（包括零命中）进行BH校正，再按最小命中数筛选展示。最小命中数设为0可导出所有检验条目。

## 本次参数与数据质量
- 有效目标基因：{diagnostics['target_count']}
- 背景基因：{diagnostics['background_count']}
- 注释关系：{diagnostics['annotation_count']}
- 被检验条目：{diagnostics['term_count']}
- 最小命中基因数：{minimum_overlap}
- FDR 阈值：{fdr_limit}
- 因不在背景集中而忽略的目标基因：{diagnostics['ignored_target_count']}

## 解释边界
显著富集表示目标基因在某功能条目中出现得比随机期望更频繁，不等于该通路已被实验验证，也不证明因果关系。结论会受到背景基因、注释数据库版本、基因标识标准化和筛选阈值影响。
"""


def page_enrichment_lab(cfg, project_id, embedded=False):
    st.subheader("上传数据并计算富集" if embedded else "GO/KEGG 在线富集实验室")
    if not embedded:
        st.info("输入准备见侧栏「核心学习与实训 → 04 GO/KEGG富集」；本站当前执行ORA。")
    st.write(
        "上传目标基因、背景基因和“基因—条目”注释表，平台将现场执行"
        "超几何检验与 FDR 校正，而不是展示预先算好的结果。"
    )

    learn_tab, prepare_tab, calculate_tab, interpret_tab = (
        (st.expander("查看计算原理"), st.expander("输入格式与示例下载"), st.container(), st.expander("怎样解读结果"))
        if embedded else st.tabs(
        ["① 学原理", "② 备数据", "③ 做分析", "④ 读结果"]
    ))

    with learn_tab:
        st.markdown(
            """
### 这个实验在回答什么？

假设背景中有一部分基因属于某条通路。如果你的目标基因中属于该通路的
基因明显多于随机抽样的期望数量，就称为“过度富集”。

**目标基因**可以是共同靶点、差异基因或模型筛选基因；**背景基因**应是
本次研究中实际可能被选中的全部基因，而不是随意使用全基因组。

计算步骤：统一基因符号并去重；限制在背景集内；执行单侧超几何检验；
使用 Benjamini–Hochberg 方法校正；结合 FDR、富集倍数和命中数解释。
            """
        )
        st.warning(
            "不要只看 P 值。背景集选择不当、少量基因造成的极高富集倍数，"
            "以及数据库注释偏倚，都可能导致误读。"
        )

    templates = example_files()
    with prepare_tab:
        st.markdown(
            """
### 三个输入文件

- `target_genes.csv`：至少包含 `gene_symbol`，每行一个目标基因。
- `background_genes.csv`：至少包含 `gene_symbol`；可不上传，此时以注释集中全部基因为背景。
- `term_annotations.csv`：每行是一条基因—功能条目关系，必须包含 `gene_symbol`、`term_id`、`term_name`，可选 `category`。

注释表可以来自经过许可的 GO/KEGG 注释资源。提交作业时应记录数据库名称、版本、下载日期和物种。
            """
        )
        st.caption("内置注释为缩小的教学示例，含TRAIN占位条目，不是完整或经版本核验的GO/KEGG注释库。当前基因处理按人类符号转大写；其他物种正式分析请先核对ID处理规则。")
        download_columns = st.columns(3)
        for column, (name, content) in zip(download_columns, templates.items()):
            column.download_button(
                f"下载 {name}", content, file_name=name, mime="text/csv",
                key=f"download_{name}", use_container_width=True,
            )
        st.dataframe(
            pd.read_csv(BytesIO(templates["term_annotations.csv"])),
            use_container_width=True, hide_index=True,
        )

    with calculate_tab:
        use_example = st.toggle(
            "使用内置教学示例", value=True,
            help="首次学习建议打开；关闭后可上传自己的三个文件。",
        )
        target_file = background_file = annotation_file = None
        pasted_genes = ""
        if not use_example:
            left, middle, right = st.columns(3)
            target_file = left.file_uploader(
                "目标基因表", type=["csv", "tsv", "txt", "xlsx", "xls"],
                key="ora_target_file",
            )
            background_file = middle.file_uploader(
                "背景基因表（可选）", type=["csv", "tsv", "txt", "xlsx", "xls"],
                key="ora_background_file",
            )
            annotation_file = right.file_uploader(
                "GO/KEGG 注释表", type=["csv", "tsv", "txt", "xlsx", "xls"],
                key="ora_annotation_file",
            )
            pasted_genes = st.text_area(
                "也可以直接粘贴目标基因", placeholder="TNF, IL6, AKT1\nMAPK1",
                help="上传目标基因表和粘贴内容同时存在时，平台会合并并去重。",
            )

        parameter1, parameter2 = st.columns(2)
        minimum_overlap = parameter1.number_input(
            "最小命中基因数", min_value=0, max_value=20, value=2, step=1,
            help="仅影响展示，不缩小BH校正范围；设为0保留全部检验条目。"
        )
        fdr_limit = parameter2.number_input(
            "FDR 阈值", min_value=0.001, max_value=1.0, value=0.05,
            step=0.01, format="%.3f"
        )

        if st.button("开始计算富集分析", type="primary", use_container_width=True):
            try:
                if use_example:
                    target_data = pd.read_csv(BytesIO(templates["target_genes.csv"]))
                    background_data = pd.read_csv(BytesIO(templates["background_genes.csv"]))
                    annotations = pd.read_csv(BytesIO(templates["term_annotations.csv"]))
                    target_genes, _ = extract_genes(target_data)
                    background_genes, _ = extract_genes(background_data)
                else:
                    target_genes = split_gene_text(pasted_genes)
                    if target_file is not None:
                        uploaded_targets, _ = extract_genes(read_table(target_file))
                        target_genes |= uploaded_targets
                    if not target_genes:
                        raise ValueError("请上传目标基因表或粘贴目标基因。")
                    if annotation_file is None:
                        raise ValueError("请上传 GO/KEGG 注释表。")
                    annotations = read_table(annotation_file)
                    background_genes = None
                    if background_file is not None:
                        background_genes, _ = extract_genes(read_table(background_file))

                result, diagnostics = run_over_representation(
                    target_genes, annotations, background_genes, int(minimum_overlap)
                )
                st.session_state["ora_result"] = result
                st.session_state["ora_diagnostics"] = diagnostics
                st.session_state["ora_parameters"] = {
                    "minimum_overlap": int(minimum_overlap),
                    "fdr_limit": float(fdr_limit),
                }
            except Exception as error:
                for result_key in ["ora_result", "ora_diagnostics", "ora_parameters"]:
                    st.session_state.pop(result_key, None)
                st.error(f"计算未完成：{error}")

        result = st.session_state.get("ora_result")
        diagnostics = st.session_state.get("ora_diagnostics")
        parameters = st.session_state.get("ora_parameters", {})
        if result is not None and diagnostics is not None:
            st.success("计算完成。以下结果由本次输入数据现场生成。")
            metric_columns = st.columns(4)
            metric_columns[0].metric("有效目标基因", diagnostics["target_count"])
            metric_columns[1].metric(
                "有效背景基因",
                diagnostics["background_count"],
                help="上传背景基因与注释表基因的交集，只有这些基因进入统计总体。",
            )
            metric_columns[2].metric("检验条目", diagnostics["term_count"])
            metric_columns[3].metric("忽略目标", diagnostics["ignored_target_count"])

            active_fdr = parameters.get("fdr_limit", float(fdr_limit))
            significant = result[result["p_adjust"] <= active_fdr].copy()
            st.write(f"FDR ≤ {active_fdr:g} 的条目：**{len(significant)}**")
            if diagnostics["ignored_target_count"]:
                with st.expander("查看未进入计算的目标基因"):
                    st.write("、".join(diagnostics["ignored_targets"]))

            display_result = significant if not significant.empty else result
            if display_result.empty:
                st.warning("没有条目达到最小命中数，请检查注释集或降低该参数。")
            else:
                st.dataframe(display_result, use_container_width=True, hide_index=True)
                _plot_bubble(display_result, "在线计算富集气泡图")
                st.download_button(
                    "下载命中数筛选后的全部结果（CSV）",
                    result.to_csv(index=False).encode("utf-8-sig"),
                    "online_enrichment_results.csv", "text/csv",
                    use_container_width=True,
                )
                method_text = _enrichment_method_text(
                    diagnostics, parameters.get("minimum_overlap", int(minimum_overlap)),
                    active_fdr,
                )
                st.download_button(
                    "下载实验方法与质量记录（Markdown）",
                    method_text.encode("utf-8-sig"),
                    "enrichment_experiment_record.md", "text/markdown",
                    use_container_width=True,
                )

    with interpret_tab:
        st.markdown(
            """
### 结果解读清单

1. **先看输入质量**：有多少目标基因被背景集接纳？忽略了哪些基因？
2. **再看 FDR**：`p_adjust` 是多重检验校正后的显著性。
3. **同时看命中数**：只有一两个基因支撑的条目需要谨慎。
4. **查看富集倍数**：大于 1 表示高于随机期望，小条目可能夸大倍数。
5. **检查命中基因**：相关通路是否由同一批高频基因反复驱动？
6. **回到研究问题**：统计富集不是机制验证，应结合文献与独立实验。

### 课堂思考题

- 更换为实验实际检测到的背景基因后，结果如何变化？
- 把最小命中数从 1 调到 3，哪些条目消失？为什么？
- 为什么不能把最小的 P 值直接称为“最重要的通路”？
            """
        )


def page_enrichment(cfg, project_id):
    st.header("GO / KEGG 富集学习与分析")
    lab_tab, case_tab = st.tabs(["教学实验：上传并计算", "五积散示范结果"])
    with lab_tab:
        page_enrichment_lab(cfg, project_id)
    with case_tab:
        page_enrichment_results(cfg, project_id)


def _plot_bubble(df, title):
    import numpy as np
    import plotly.express as px

    needed = {"gene_ratio", "p_adjust", "term_name"}

    if not needed.issubset(df.columns):
        st.warning("富集结果缺少绘图所需字段。")
        return

    plot_df = df.copy()

    plot_df["gene_ratio"] = pd.to_numeric(
        plot_df["gene_ratio"],
        errors="coerce",
    )

    plot_df["p_adjust"] = pd.to_numeric(
        plot_df["p_adjust"],
        errors="coerce",
    )

    plot_df = plot_df.dropna(
        subset=["gene_ratio", "p_adjust", "term_name"]
    )

    top_n = st.slider(
        f"{title}显示条目数",
        min_value=10,
        max_value=50,
        value=20,
        step=5,
        key=f"{title}_top_n",
    )

    plot_df = (
        plot_df
        .sort_values("p_adjust", ascending=True)
        .head(top_n)
        .copy()
    )

    plot_df["显著性"] = -np.log10(
        plot_df["p_adjust"].clip(lower=1e-300)
    )

    size_column = (
        "count"
        if "count" in plot_df.columns
        else None
    )

    fig = px.scatter(
        plot_df,
        x="gene_ratio",
        y="term_name",
        size=size_column,
        color="显著性",
        color_continuous_scale="YlOrRd",
        labels={
            "gene_ratio": "GeneRatio",
            "term_name": "富集条目",
            "显著性": "-log10(校正后P值)",
        },
        title=title,
        height=max(520, top_n * 30),
    )

    fig.update_layout(
        yaxis={"categoryorder": "total ascending"},
        margin=dict(l=20, r=20, t=60, b=20),
    )

    st.plotly_chart(fig, use_container_width=True)
# ---------------------------------------------------------------------------
# 页面：转录组验证
# ---------------------------------------------------------------------------
def page_transcriptome(cfg, project_id):
    import numpy as np
    import plotly.express as px

    st.header("转录组证据验证")
    st.info("本页解读已有项目结果。新数据请进入「进阶证据 → 转录组数据与设计」第5步：可上传标准化log2表达，做独立两组Welch差异分析并衔接富集。原始计数、配对和批次模型按课程外部路线执行。")

    files = cfg.get("data_files", {})

    degs = load_csv(
        project_path(
            project_id,
            files.get("deg_results", "")
        )
    )

    samples = load_csv(
        project_path(
            project_id,
            files.get("transcriptome_samples", "")
        )
    )

    overlap = load_csv(
        project_path(
            project_id,
            files.get("transcriptome_target_overlap", "")
        )
    )

    if degs is None:
        st.info(
            "未找到转录组差异表达结果，"
            "请检查 project.yaml 中的 deg_results。"
        )
        return

    degs = degs.copy()

    required_columns = {
        "gene_symbol",
        "logFC",
        "adjusted_p_value",
        "regulation",
    }

    if not required_columns.issubset(degs.columns):
        st.error(
            "转录组结果缺少必要字段："
            "gene_symbol、logFC、adjusted_p_value、regulation。"
        )
        return

    degs["logFC"] = pd.to_numeric(
        degs["logFC"],
        errors="coerce"
    )

    degs["adjusted_p_value"] = pd.to_numeric(
        degs["adjusted_p_value"],
        errors="coerce"
    )

    degs = degs.dropna(
        subset=[
            "gene_symbol",
            "logFC",
            "adjusted_p_value",
        ]
    )

    sample_count = len(samples) if samples is not None else 0
    control_count = 0
    uc_count = 0

    if samples is not None and "group" in samples.columns:
        control_count = int(
            (samples["group"] == "Control").sum()
        )
        uc_count = int(
            (samples["group"] == "UC").sum()
        )

    up_count = int(
        (degs["regulation"] == "Up").sum()
    )

    down_count = int(
        (degs["regulation"] == "Down").sum()
    )

    if overlap is not None and "gene_symbol" in overlap.columns:
        supported_count = overlap["gene_symbol"].nunique()
    else:
        supported_count = 0

    st.subheader("GSE38713 数据概览")

    metric1, metric2, metric3, metric4 = st.columns(4)

    metric1.metric(
        "分析样本",
        f"{sample_count:,}",
        help=f"正常对照 {control_count}，UC患者 {uc_count}"
    )

    metric2.metric(
        "显著差异基因",
        f"{len(degs):,}"
    )

    metric3.metric(
        "上调 / 下调",
        f"{up_count:,} / {down_count:,}"
    )

    metric4.metric(
        "多证据支持靶点",
        f"{supported_count:,}"
    )

    st.caption(
        "比较方向：UC vs Control。正 logFC 表示在UC中上调，"
        "负 logFC 表示在UC中下调。显著性标准为"
        "校正后P值 < 0.05且 |logFC| > 0.58。"
    )

    overview_tab, deg_tab, evidence_tab = st.tabs(
        [
            "样本与表达概览",
            "差异表达基因",
            "多证据共同靶点",
        ]
    )

    # --------------------------------------------------
    # 样本与表达概览
    # --------------------------------------------------
    with overview_tab:
        st.subheader("样本分组")

        if samples is None:
            st.warning("未找到转录组样本分组文件。")
        else:
            if "group" in samples.columns:
                group_counts = (
                    samples["group"]
                    .value_counts()
                    .rename_axis("group")
                    .reset_index(name="sample_count")
                )

                group_counts["group"] = group_counts[
                    "group"
                ].replace(
                    {
                        "Control": "正常对照",
                        "UC": "溃疡性结肠炎",
                    }
                )

                sample_fig = px.bar(
                    group_counts,
                    x="group",
                    y="sample_count",
                    color="group",
                    text="sample_count",
                    color_discrete_map={
                        "正常对照": "#2E8B57",
                        "溃疡性结肠炎": "#C0392B",
                    },
                    labels={
                        "group": "样本分组",
                        "sample_count": "样本数量",
                    },
                    title="GSE38713样本构成",
                    height=430,
                )

                sample_fig.update_traces(
                    textposition="outside"
                )

                sample_fig.update_layout(
                    showlegend=False
                )

                st.plotly_chart(
                    sample_fig,
                    use_container_width=True
                )

            with st.expander("查看样本信息"):
                st.dataframe(
                    samples,
                    use_container_width=True,
                    hide_index=True
                )

        st.info(
            "本模块使用GSE38713结肠活检表达数据，"
            "比较活动期溃疡性结肠炎患者与非炎症正常对照。"
        )

    # --------------------------------------------------
    # 差异表达基因
    # --------------------------------------------------
    with deg_tab:
        st.subheader("差异基因筛选")

        filter1, filter2, filter3 = st.columns(3)

        search_gene = filter1.text_input(
            "搜索基因",
            placeholder="例如：TNF、IL6、CTSB",
            key="transcriptome_gene_search"
        )

        regulation_option = filter2.selectbox(
            "调控方向",
            ["全部", "Up", "Down"],
            key="transcriptome_regulation"
        )

        maximum_fdr = filter3.number_input(
            "校正后P值上限",
            min_value=0.0,
            max_value=0.05,
            value=0.05,
            step=0.001,
            format="%.3f",
            key="transcriptome_fdr"
        )

        filtered_degs = degs[
            degs["adjusted_p_value"] <= maximum_fdr
        ].copy()

        if regulation_option != "全部":
            filtered_degs = filtered_degs[
                filtered_degs["regulation"]
                == regulation_option
            ]

        if search_gene:
            filtered_degs = filtered_degs[
                filtered_degs["gene_symbol"]
                .astype(str)
                .str.contains(
                    search_gene,
                    case=False,
                    regex=False,
                    na=False
                )
            ]

        result1, result2, result3 = st.columns(3)

        result1.metric(
            "当前结果",
            f"{len(filtered_degs):,}"
        )

        result2.metric(
            "当前上调",
            f"{(filtered_degs['regulation'] == 'Up').sum():,}"
        )

        result3.metric(
            "当前下调",
            f"{(filtered_degs['regulation'] == 'Down').sum():,}"
        )

        plot_degs = filtered_degs.copy()

        if not plot_degs.empty:
            plot_degs["negative_log10_fdr"] = -np.log10(
                plot_degs[
                    "adjusted_p_value"
                ].clip(lower=1e-300)
            )

            volcano_fig = px.scatter(
                plot_degs,
                x="logFC",
                y="negative_log10_fdr",
                color="regulation",
                color_discrete_map={
                    "Up": "#C0392B",
                    "Down": "#2874A6",
                },
                hover_name="gene_symbol",
                hover_data={
                    "logFC": ":.3f",
                    "adjusted_p_value": ":.3e",
                    "negative_log10_fdr": False,
                },
                labels={
                    "logFC": "logFC（UC vs Control）",
                    "negative_log10_fdr":
                        "-log10（校正后P值）",
                    "regulation": "调控方向",
                },
                title="显著差异基因分布",
                height=570,
            )

            volcano_fig.add_vline(
                x=0.58,
                line_dash="dash",
                line_color="#777777"
            )

            volcano_fig.add_vline(
                x=-0.58,
                line_dash="dash",
                line_color="#777777"
            )

            st.plotly_chart(
                volcano_fig,
                use_container_width=True
            )

            st.caption(
                "注意：当前图仅展示已经达到显著性标准的基因，"
                "不包含未达到显著性的背景基因，"
                "因此属于显著差异基因分布图。"
            )

            top_genes = (
                plot_degs
                .assign(
                    absolute_logfc=plot_degs[
                        "logFC"
                    ].abs()
                )
                .sort_values(
                    "absolute_logfc",
                    ascending=False
                )
                .head(20)
                .sort_values("logFC")
            )

            top_fig = px.bar(
                top_genes,
                x="logFC",
                y="gene_symbol",
                orientation="h",
                color="regulation",
                color_discrete_map={
                    "Up": "#C0392B",
                    "Down": "#2874A6",
                },
                labels={
                    "logFC": "logFC",
                    "gene_symbol": "基因",
                    "regulation": "调控方向",
                },
                title="|logFC|最高的20个差异基因",
                height=600,
            )

            st.plotly_chart(
                top_fig,
                use_container_width=True
            )
        else:
            st.warning("当前筛选条件下没有差异基因。")

        deg_csv = filtered_degs.to_csv(
            index=False
        ).encode("utf-8-sig")

        st.download_button(
            "下载当前差异基因结果（CSV）",
            data=deg_csv,
            file_name="filtered_transcriptome_degs.csv",
            mime="text/csv",
            key="download_transcriptome_degs"
        )

        st.dataframe(
            filtered_degs,
            use_container_width=True,
            hide_index=True
        )

    # --------------------------------------------------
    # 多证据共同靶点
    # --------------------------------------------------
    with evidence_tab:
        st.subheader(
            "网络药理学与转录组共同支持的候选靶点"
        )

        if overlap is None or overlap.empty:
            st.warning(
                "没有找到转录组与网络药理学交集结果。"
            )
        else:
            overlap_filtered = overlap.copy()

            evidence1, evidence2 = st.columns(2)

            overlap_search = evidence1.text_input(
                "搜索多证据靶点",
                placeholder="例如：EGFR、CTSB、CXCL1",
                key="transcriptome_overlap_search"
            )

            overlap_regulation = evidence2.selectbox(
                "表达方向",
                ["全部", "Up", "Down"],
                key="transcriptome_overlap_regulation"
            )

            if overlap_search:
                overlap_filtered = overlap_filtered[
                    overlap_filtered["gene_symbol"]
                    .astype(str)
                    .str.contains(
                        overlap_search,
                        case=False,
                        regex=False,
                        na=False
                    )
                ]

            if (
                overlap_regulation != "全部"
                and "regulation" in overlap_filtered.columns
            ):
                overlap_filtered = overlap_filtered[
                    overlap_filtered["regulation"]
                    == overlap_regulation
                ]

            st.metric(
                "当前多证据共同靶点",
                f"{overlap_filtered['gene_symbol'].nunique():,}"
            )

            if (
                not overlap_filtered.empty
                and "logFC" in overlap_filtered.columns
            ):
                evidence_plot = (
                    overlap_filtered
                    .assign(
                        absolute_logfc=lambda data:
                            pd.to_numeric(
                                data["logFC"],
                                errors="coerce"
                            ).abs()
                    )
                    .sort_values(
                        "absolute_logfc",
                        ascending=False
                    )
                    .head(25)
                    .sort_values("logFC")
                )

                evidence_fig = px.bar(
                    evidence_plot,
                    x="logFC",
                    y="gene_symbol",
                    orientation="h",
                    color="regulation",
                    color_discrete_map={
                        "Up": "#C0392B",
                        "Down": "#2874A6",
                    },
                    labels={
                        "logFC": "logFC",
                        "gene_symbol": "共同靶点",
                        "regulation": "表达方向",
                    },
                    title="多证据共同靶点表达变化 Top 25",
                    height=680,
                )

                st.plotly_chart(
                    evidence_fig,
                    use_container_width=True
                )

            overlap_csv = overlap_filtered.to_csv(
                index=False
            ).encode("utf-8-sig")

            st.download_button(
                "下载多证据共同靶点（CSV）",
                data=overlap_csv,
                file_name=(
                    "transcriptome_supported_targets.csv"
                ),
                mime="text/csv",
                key="download_transcriptome_overlap"
            )

            st.dataframe(
                overlap_filtered,
                use_container_width=True,
                hide_index=True
            )

        st.warning(
            "网络药理学共同靶点与转录组差异基因的交集，"
            "只能提高候选靶点的证据等级，不能直接证明五积散"
            "对溃疡性结肠炎具有治疗作用或因果机制。"
        )

# ---------------------------------------------------------------------------
# 页面：机器学习 / defMR / 其他组学
# ---------------------------------------------------------------------------
# ---------------------------------------------------------------------------
# 页面：机器学习结果
# ---------------------------------------------------------------------------
def page_machine_learning(cfg, project_id):
    import plotly.express as px

    st.header("机器学习模型评估")
    st.info("本页解读已有项目文件。新数据请进入「进阶证据 → 机器学习与评估」实际操作区，上传表达和分组，公平比较Logistic、Elastic Net、随机森林与SVM并评估。")

    files = cfg.get("data_files", {})

    def get_ml_data(file_key):
        return load_csv(
            project_path(
                project_id,
                files.get(file_key, "")
            )
        )

    metrics = get_ml_data("ml_metrics")
    auc_data = get_ml_data("ml_auc")
    features = get_ml_data("ml_features")
    consensus = get_ml_data("ml_feature_consensus")
    quality_flags = get_ml_data("ml_quality_flags")

    if metrics is None:
        st.info(
            "未找到机器学习性能结果，"
            "请检查 project.yaml 中的 ml_metrics。"
        )
        return

    metrics = metrics.copy()

    model_count = (
        metrics["Algorithm"].nunique()
        if "Algorithm" in metrics.columns
        else 0
    )

    best_external_auc = 0.0

    if (
        auc_data is not None
        and {"evaluation_set", "AUC"}.issubset(
            auc_data.columns
        )
    ):
        auc_data = auc_data.copy()
        auc_data["AUC"] = pd.to_numeric(
            auc_data["AUC"],
            errors="coerce"
        )

        external_auc = auc_data[
            auc_data["evaluation_set"]
            == "External GSE38713"
        ]

        if not external_auc.empty:
            best_external_auc = (
                external_auc["AUC"].max()
            )

    multi_model_gene_count = 0

    if (
        consensus is not None
        and "model_count" in consensus.columns
    ):
        consensus = consensus.copy()
        consensus["model_count"] = pd.to_numeric(
            consensus["model_count"],
            errors="coerce"
        )

        multi_model_gene_count = int(
            (consensus["model_count"] >= 2).sum()
        )

    quality_flag_count = (
        len(quality_flags)
        if quality_flags is not None
        else 0
    )

    metric1, metric2, metric3, metric4 = st.columns(4)

    metric1.metric(
        "比较模型",
        f"{model_count}"
    )

    metric2.metric(
        "最佳外部验证AUC",
        f"{best_external_auc:.3f}"
    )

    metric3.metric(
        "多模型支持基因",
        f"{multi_model_gene_count}"
    )

    metric4.metric(
        "质量风险提示",
        f"{quality_flag_count}"
    )

    st.caption(
        "本模块展示已有机器学习模型的训练、内部验证和"
        "GSE38713外部验证结果。模型性能仅用于科研评估，"
        "不能直接作为临床诊断工具。"
    )

    performance_tab, auc_tab, feature_tab, quality_tab = st.tabs(
        [
            "模型性能",
            "跨数据集AUC",
            "特征基因",
            "质量评估",
        ]
    )

    # --------------------------------------------------
    # 模型性能
    # --------------------------------------------------
    with performance_tab:
        st.subheader("训练与内部验证性能")

        if "evaluation_set" not in metrics.columns:
            st.warning("性能结果缺少 evaluation_set 字段。")
        else:
            evaluation_options = (
                metrics["evaluation_set"]
                .dropna()
                .astype(str)
                .unique()
                .tolist()
            )

            selected_evaluation = st.selectbox(
                "选择评估数据集",
                evaluation_options,
                key="ml_evaluation_set"
            )

            selected_metrics = metrics[
                metrics["evaluation_set"]
                == selected_evaluation
            ].copy()

            metric_columns = [
                column
                for column in [
                    "Accuracy",
                    "AUC_ROC",
                    "Sensitivity",
                    "Specificity",
                    "Precision",
                    "F1",
                    "MCC",
                ]
                if column in selected_metrics.columns
            ]

            for column in metric_columns:
                selected_metrics[column] = pd.to_numeric(
                    selected_metrics[column],
                    errors="coerce"
                )

            if metric_columns:
                performance_long = selected_metrics.melt(
                    id_vars=["Algorithm"],
                    value_vars=metric_columns,
                    var_name="metric",
                    value_name="score"
                )

                performance_fig = px.bar(
                    performance_long,
                    x="Algorithm",
                    y="score",
                    color="metric",
                    barmode="group",
                    labels={
                        "Algorithm": "模型",
                        "score": "性能得分",
                        "metric": "评价指标",
                    },
                    title=(
                        f"{selected_evaluation}模型性能比较"
                    ),
                    height=600,
                )

                performance_fig.update_yaxes(
                    range=[0, 1.05]
                )

                st.plotly_chart(
                    performance_fig,
                    use_container_width=True
                )

            st.dataframe(
                selected_metrics,
                use_container_width=True,
                hide_index=True
            )

            metrics_csv = selected_metrics.to_csv(
                index=False
            ).encode("utf-8-sig")

            st.download_button(
                "下载当前模型性能（CSV）",
                data=metrics_csv,
                file_name="machine_learning_metrics.csv",
                mime="text/csv",
                key="download_ml_metrics"
            )

    # --------------------------------------------------
    # 跨数据集AUC
    # --------------------------------------------------
    with auc_tab:
        st.subheader("训练、内部验证与外部验证AUC")

        if auc_data is None:
            st.warning("未找到跨数据集AUC结果。")
        else:
            auc_plot = auc_data.dropna(
                subset=["Algorithm", "evaluation_set", "AUC"]
            ).copy()

            auc_fig = px.line(
                auc_plot,
                x="evaluation_set",
                y="AUC",
                color="Algorithm",
                markers=True,
                category_orders={
                    "evaluation_set": [
                        "Training",
                        "Internal validation",
                        "External GSE38713",
                    ]
                },
                labels={
                    "evaluation_set": "评估数据集",
                    "AUC": "AUC",
                    "Algorithm": "模型",
                },
                title="模型跨数据集AUC变化",
                height=570,
            )

            auc_fig.update_traces(
                marker={"size": 10}
            )

            auc_fig.update_yaxes(
                range=[0, 1.05]
            )

            st.plotly_chart(
                auc_fig,
                use_container_width=True
            )

            auc_table = auc_plot.pivot_table(
                index="Algorithm",
                columns="evaluation_set",
                values="AUC",
                aggfunc="first"
            ).reset_index()

            st.dataframe(
                auc_table,
                use_container_width=True,
                hide_index=True
            )

            auc_csv = auc_plot.to_csv(
                index=False
            ).encode("utf-8-sig")

            st.download_button(
                "下载跨数据集AUC（CSV）",
                data=auc_csv,
                file_name="model_dataset_auc.csv",
                mime="text/csv",
                key="download_ml_auc"
            )

            st.warning(
                "多项AUC达到1.0。异常理想的性能可能与样本量、"
                "数据划分、批次效应、特征筛选顺序或数据泄漏有关，"
                "需要使用更多完全独立的队列复现。"
            )

    # --------------------------------------------------
    # 特征基因
    # --------------------------------------------------
    with feature_tab:
        st.subheader("多模型特征基因")

        if consensus is None:
            st.warning("未找到多模型特征基因汇总。")
        else:
            consensus = consensus.copy()

            consensus[
                "mean_normalized_importance"
            ] = pd.to_numeric(
                consensus["mean_normalized_importance"],
                errors="coerce"
            )

            top_consensus = (
                consensus
                .dropna(
                    subset=[
                        "gene_symbol",
                        "mean_normalized_importance",
                    ]
                )
                .sort_values(
                    [
                        "model_count",
                        "mean_normalized_importance",
                    ],
                    ascending=[False, False]
                )
                .head(25)
                .sort_values(
                    "mean_normalized_importance"
                )
            )

            consensus_fig = px.bar(
                top_consensus,
                x="mean_normalized_importance",
                y="gene_symbol",
                orientation="h",
                color="model_count",
                color_continuous_scale="YlGn",
                hover_data=["models", "best_rank"],
                labels={
                    "mean_normalized_importance":
                        "平均标准化重要性",
                    "gene_symbol": "特征基因",
                    "model_count": "支持模型数",
                    "models": "支持模型",
                    "best_rank": "最佳排名",
                },
                title="多模型共同支持特征基因 Top 25",
                height=720,
            )

            st.plotly_chart(
                consensus_fig,
                use_container_width=True
            )

            consensus_search = st.text_input(
                "搜索共同特征基因",
                placeholder="例如：TIMP1、SERPINA1、IL1B",
                key="ml_consensus_search"
            )

            consensus_filtered = consensus.copy()

            if consensus_search:
                consensus_filtered = consensus_filtered[
                    consensus_filtered["gene_symbol"]
                    .astype(str)
                    .str.contains(
                        consensus_search,
                        case=False,
                        regex=False,
                        na=False
                    )
                ]

            st.dataframe(
                consensus_filtered,
                use_container_width=True,
                hide_index=True
            )

            consensus_csv = consensus_filtered.to_csv(
                index=False
            ).encode("utf-8-sig")

            st.download_button(
                "下载多模型特征基因（CSV）",
                data=consensus_csv,
                file_name="ml_feature_consensus.csv",
                mime="text/csv",
                key="download_ml_consensus"
            )

        st.divider()
        st.subheader("单模型特征重要性")

        if features is None:
            st.warning("未找到单模型特征重要性数据。")
        else:
            features = features.copy()

            algorithm_options = (
                features["Algorithm"]
                .dropna()
                .astype(str)
                .unique()
                .tolist()
            )

            selected_algorithm = st.selectbox(
                "选择模型",
                algorithm_options,
                key="ml_feature_algorithm"
            )

            selected_features = features[
                features["Algorithm"]
                == selected_algorithm
            ].copy()

            selected_features[
                "normalized_importance"
            ] = pd.to_numeric(
                selected_features["normalized_importance"],
                errors="coerce"
            )

            top_features = (
                selected_features
                .sort_values(
                    "normalized_importance",
                    ascending=False
                )
                .head(20)
                .sort_values(
                    "normalized_importance"
                )
            )

            feature_fig = px.bar(
                top_features,
                x="normalized_importance",
                y="gene_symbol",
                orientation="h",
                color="normalized_importance",
                color_continuous_scale="Blues",
                labels={
                    "normalized_importance":
                        "模型内标准化重要性",
                    "gene_symbol": "特征基因",
                },
                title=(
                    f"{selected_algorithm}特征重要性 Top 20"
                ),
                height=620,
            )

            feature_fig.update_layout(
                coloraxis_showscale=False
            )

            st.plotly_chart(
                feature_fig,
                use_container_width=True
            )

            st.dataframe(
                selected_features,
                use_container_width=True,
                hide_index=True
            )

    # --------------------------------------------------
    # 质量评估
    # --------------------------------------------------
    with quality_tab:
        st.subheader("模型结果质量与风险检查")

        if quality_flags is None or quality_flags.empty:
            st.success("当前未识别到明显质量风险。")
        else:
            high_count = int(
                (
                    quality_flags["severity"]
                    == "High"
                ).sum()
            )

            medium_count = int(
                (
                    quality_flags["severity"]
                    == "Medium"
                ).sum()
            )

            quality1, quality2 = st.columns(2)

            quality1.metric(
                "高风险提示",
                high_count
            )

            quality2.metric(
                "中等风险提示",
                medium_count
            )

            if high_count > 0:
                st.error(
                    "发现高风险性能问题。请重点检查特异度、"
                    "分类阈值和混淆矩阵。"
                )

            st.dataframe(
                quality_flags,
                use_container_width=True,
                hide_index=True
            )

            quality_csv = quality_flags.to_csv(
                index=False
            ).encode("utf-8-sig")

            st.download_button(
                "下载质量风险记录（CSV）",
                data=quality_csv,
                file_name="ml_quality_flags.csv",
                mime="text/csv",
                key="download_ml_quality_flags"
            )

        st.info(
            "建议后续检查：训练集与测试集是否严格独立、"
            "特征筛选是否只在训练集内完成、是否存在批次效应、"
            "是否报告混淆矩阵和置信区间，以及模型能否在新的"
            "独立人群中复现。"
        )

        st.warning(
            "机器学习模型发现的是预测关联，而不是治疗因果机制。"
            "高AUC不能直接证明某个基因是五积散治疗"
            "溃疡性结肠炎的作用靶点。"
        )
# ---------------------------------------------------------------------------
# 页面：孟德尔随机化
# ---------------------------------------------------------------------------
def page_mendelian_randomization(cfg, project_id):
    import numpy as np
    import plotly.express as px
    import plotly.graph_objects as go

    st.header("孟德尔随机化证据评估")
    st.info("本页解读既有项目结果。完整资料与假设教学见侧栏「进阶证据 → 孟德尔随机化」。公式示意不等于通用MR计算服务。")

    files = cfg.get("data_files", {})

    def get_mr_data(file_key):
        return load_csv(
            project_path(
                project_id,
                files.get(file_key, "")
            )
        )

    mr_results = get_mr_data("mr_results")
    method_results = get_mr_data("mr_method_results")
    sensitivity = get_mr_data("mr_sensitivity")
    comparison = get_mr_data("mr_strategy_comparison")
    presso = get_mr_data("mr_presso")
    evidence = get_mr_data("mr_evidence_summary")

    if mr_results is None:
        st.info(
            "未找到孟德尔随机化结果，"
            "请检查project.yaml中的mr_results。"
        )
        return

    mr_results = mr_results.copy()

    numeric_columns = [
        "valid_snp_count",
        "odds_ratio",
        "ci_lower",
        "ci_upper",
        "p_value",
        "fdr_adjusted_p",
    ]

    for column in numeric_columns:
        if column in mr_results.columns:
            mr_results[column] = pd.to_numeric(
                mr_results[column],
                errors="coerce"
            )

    def as_boolean(series):
        """兼容CSV读取后的布尔值和文字值。"""
        if series.dtype == bool:
            return series.fillna(False)

        return (
            series
            .astype(str)
            .str.strip()
            .str.lower()
            .isin(["true", "1", "yes"])
        )

    full_results = mr_results[
        mr_results["instrument_strategy"]
        == "All available instruments"
    ].copy()

    analyzed_genes = (
        mr_results["gene_symbol"].nunique()
    )

    full_significant = int(
        as_boolean(
            full_results["nominal_significant"]
        ).sum()
    )

    direction_consistent_count = 0

    if (
        comparison is not None
        and "direction_consistent"
        in comparison.columns
    ):
        direction_consistent_count = int(
            as_boolean(
                comparison["direction_consistent"]
            ).sum()
        )

    pleiotropy_count = 0

    if (
        sensitivity is not None
        and "pleiotropy_flag"
        in sensitivity.columns
    ):
        pleiotropy_count = int(
            as_boolean(
                sensitivity["pleiotropy_flag"]
            ).sum()
        )

    metric1, metric2, metric3, metric4 = st.columns(4)

    metric1.metric(
        "分析基因",
        analyzed_genes
    )

    metric2.metric(
        "全SNP名义显著",
        full_significant
    )

    metric3.metric(
        "策略方向一致",
        direction_consistent_count
    )

    metric4.metric(
        "多效性风险",
        pleiotropy_count
    )

    st.caption(
        "结局：溃疡性结肠炎（GWAS ID：ieu-a-32）。"
        "主要估计方法为逆方差加权法（IVW）。"
        "P < 0.05仅表示名义显著，还需要结合多重检验、"
        "异质性、水平多效性和不同工具变量策略综合判断。"
    )

    overview_tab, strategy_tab, method_tab, sensitivity_tab = (
        st.tabs(
            [
                "主要MR结果",
                "策略一致性",
                "多方法验证",
                "敏感性检查",
            ]
        )
    )

    # --------------------------------------------------
    # 主要MR结果
    # --------------------------------------------------
    with overview_tab:
        st.subheader("候选基因MR效应估计")

        strategy_options = (
            mr_results["instrument_strategy"]
            .dropna()
            .astype(str)
            .unique()
            .tolist()
        )

        strategy_labels = {
            "All available instruments":
                "全部可用工具变量",
            "Top 50 instruments":
                "Top 50工具变量",
        }

        selected_strategy = st.selectbox(
            "选择SNP筛选策略",
            strategy_options,
            format_func=lambda value:
                strategy_labels.get(value, value),
            key="mr_main_strategy"
        )

        selected_results = mr_results[
            mr_results["instrument_strategy"]
            == selected_strategy
        ].copy()

        selected_results = selected_results.dropna(
            subset=[
                "gene_symbol",
                "odds_ratio",
                "ci_lower",
                "ci_upper",
            ]
        )

        selected_results["significance_label"] = np.where(
            selected_results["p_value"] < 0.05,
            "P < 0.05",
            "P ≥ 0.05"
        )

        selected_results = selected_results.sort_values(
            "odds_ratio",
            ascending=True
        )

        point_colors = selected_results[
            "significance_label"
        ].map(
            {
                "P < 0.05": "#C0392B",
                "P ≥ 0.05": "#7F8C8D",
            }
        ).tolist()

        forest_fig = go.Figure()

        forest_fig.add_trace(
            go.Scatter(
                x=selected_results["odds_ratio"],
                y=selected_results["gene_symbol"],
                mode="markers",
                marker={
                    "size": 12,
                    "color": point_colors,
                },
                error_x={
                    "type": "data",
                    "symmetric": False,
                    "array": (
                        selected_results["ci_upper"]
                        - selected_results["odds_ratio"]
                    ),
                    "arrayminus": (
                        selected_results["odds_ratio"]
                        - selected_results["ci_lower"]
                    ),
                    "color": "#AAB7B8",
                    "thickness": 1.5,
                },
                customdata=np.column_stack(
                    [
                        selected_results["ci_lower"],
                        selected_results["ci_upper"],
                        selected_results["p_value"],
                        selected_results[
                            "valid_snp_count"
                        ],
                    ]
                ),
                hovertemplate=(
                    "<b>%{y}</b><br>"
                    "OR=%{x:.3f}<br>"
                    "95%CI=%{customdata[0]:.3f}"
                    "–%{customdata[1]:.3f}<br>"
                    "P=%{customdata[2]:.3e}<br>"
                    "有效SNP=%{customdata[3]:.0f}"
                    "<extra></extra>"
                ),
                showlegend=False,
            )
        )

        forest_fig.add_vline(
            x=1,
            line_dash="dash",
            line_color="#F1C40F",
            annotation_text="OR = 1"
        )

        forest_fig.update_xaxes(
            type="log",
            title="比值比 OR（对数坐标）"
        )

        forest_fig.update_yaxes(
            title="候选基因"
        )

        forest_fig.update_layout(
            title=(
                f"{strategy_labels.get(selected_strategy, selected_strategy)}"
                "：MR森林图"
            ),
            height=560,
            margin={
                "l": 20,
                "r": 20,
                "t": 70,
                "b": 20,
            },
        )

        st.plotly_chart(
            forest_fig,
            use_container_width=True
        )

        result1, result2, result3 = st.columns(3)

        result1.metric(
            "当前基因数",
            len(selected_results)
        )

        result2.metric(
            "名义显著",
            int(
                (
                    selected_results["p_value"]
                    < 0.05
                ).sum()
            )
        )

        result3.metric(
            "FDR显著",
            int(
                (
                    selected_results[
                        "fdr_adjusted_p"
                    ]
                    < 0.05
                ).sum()
            )
        )

        display_columns = [
            "gene_symbol",
            "valid_snp_count",
            "odds_ratio",
            "ci_lower",
            "ci_upper",
            "p_value",
            "fdr_adjusted_p",
            "direction",
            "nominal_significant",
            "fdr_significant",
        ]

        st.dataframe(
            selected_results[
                [
                    column
                    for column in display_columns
                    if column
                    in selected_results.columns
                ]
            ],
            use_container_width=True,
            hide_index=True
        )

        result_csv = selected_results.to_csv(
            index=False
        ).encode("utf-8-sig")

        st.download_button(
            "下载当前MR结果（CSV）",
            data=result_csv,
            file_name="mr_selected_results.csv",
            mime="text/csv",
            key="download_mr_results"
        )

        st.info(
            "OR > 1表示遗传预测的暴露升高与UC风险升高方向相关；"
            "OR < 1表示保护方向。该方向解释依赖暴露定义，"
            "不能直接理解为药物升高或降低该基因后一定产生相同效果。"
        )

    # --------------------------------------------------
    # 策略一致性
    # --------------------------------------------------
    with strategy_tab:
        st.subheader("不同SNP筛选策略的一致性")

        if comparison is None:
            st.warning("未找到MR策略比较结果。")
        else:
            comparison = comparison.copy()

            comparison_columns = [
                "all_available_instruments__odds_ratio",
                "top_50_instruments__odds_ratio",
                "all_available_instruments__p_value",
                "top_50_instruments__p_value",
            ]

            for column in comparison_columns:
                if column in comparison.columns:
                    comparison[column] = pd.to_numeric(
                        comparison[column],
                        errors="coerce"
                    )

            strategy_plot = pd.DataFrame(
                {
                    "gene_symbol": pd.concat(
                        [
                            comparison["gene_symbol"],
                            comparison["gene_symbol"],
                        ],
                        ignore_index=True
                    ),
                    "strategy": (
                        ["全部可用工具变量"]
                        * len(comparison)
                        + ["Top 50工具变量"]
                        * len(comparison)
                    ),
                    "odds_ratio": pd.concat(
                        [
                            comparison[
                                "all_available_instruments"
                                "__odds_ratio"
                            ],
                            comparison[
                                "top_50_instruments"
                                "__odds_ratio"
                            ],
                        ],
                        ignore_index=True
                    ),
                }
            )

            strategy_fig = px.line(
                strategy_plot,
                x="strategy",
                y="odds_ratio",
                color="gene_symbol",
                markers=True,
                labels={
                    "strategy": "SNP筛选策略",
                    "odds_ratio": "OR",
                    "gene_symbol": "基因",
                },
                title="两种工具变量策略下的OR变化",
                height=590,
                log_y=True,
            )

            strategy_fig.add_hline(
                y=1,
                line_dash="dash",
                line_color="#F1C40F"
            )

            strategy_fig.update_traces(
                marker={"size": 10}
            )

            st.plotly_chart(
                strategy_fig,
                use_container_width=True
            )

            consistent = as_boolean(
                comparison["direction_consistent"]
            )

            significance_consistent = as_boolean(
                comparison[
                    "significance_consistent"
                ]
            )

            status1, status2, status3 = st.columns(3)

            status1.metric(
                "方向一致",
                int(consistent.sum())
            )

            status2.metric(
                "显著性一致",
                int(
                    significance_consistent.sum()
                )
            )

            status3.metric(
                "方向不一致",
                int((~consistent).sum())
            )

            st.dataframe(
                comparison,
                use_container_width=True,
                hide_index=True
            )

            st.warning(
                "不同SNP筛选策略若得出不同方向或不同显著性，"
                "说明结论对工具变量选择较敏感。CASP1方向不一致；"
                "CXCL1和SERPINA1的显著性随策略改变。"
            )

    # --------------------------------------------------
    # 多方法验证
    # --------------------------------------------------
    with method_tab:
        st.subheader("不同MR估计方法比较")

        if method_results is None:
            st.warning("未找到多方法MR结果。")
        else:
            method_results = method_results.copy()

            method_numeric_columns = [
                "odds_ratio",
                "ci_lower",
                "ci_upper",
                "p_value",
                "snp_count",
            ]

            for column in method_numeric_columns:
                if column in method_results.columns:
                    method_results[column] = pd.to_numeric(
                        method_results[column],
                        errors="coerce"
                    )

            gene_options = sorted(
                method_results["gene_symbol"]
                .dropna()
                .astype(str)
                .unique()
                .tolist()
            )

            selected_gene = st.selectbox(
                "选择基因",
                gene_options,
                key="mr_method_gene"
            )

            gene_methods = method_results[
                method_results["gene_symbol"]
                == selected_gene
            ].copy()

            gene_methods = gene_methods.dropna(
                subset=[
                    "method",
                    "odds_ratio",
                    "ci_lower",
                    "ci_upper",
                ]
            ).sort_values("odds_ratio")

            method_colors = np.where(
                gene_methods["p_value"] < 0.05,
                "#C0392B",
                "#7F8C8D"
            )

            method_fig = go.Figure()

            method_fig.add_trace(
                go.Scatter(
                    x=gene_methods["odds_ratio"],
                    y=gene_methods["method"],
                    mode="markers",
                    marker={
                        "size": 12,
                        "color": method_colors,
                    },
                    error_x={
                        "type": "data",
                        "symmetric": False,
                        "array": (
                            gene_methods["ci_upper"]
                            - gene_methods["odds_ratio"]
                        ),
                        "arrayminus": (
                            gene_methods["odds_ratio"]
                            - gene_methods["ci_lower"]
                        ),
                    },
                    customdata=np.column_stack(
                        [
                            gene_methods["ci_lower"],
                            gene_methods["ci_upper"],
                            gene_methods["p_value"],
                        ]
                    ),
                    hovertemplate=(
                        "<b>%{y}</b><br>"
                        "OR=%{x:.3f}<br>"
                        "95%CI=%{customdata[0]:.3f}"
                        "–%{customdata[1]:.3f}<br>"
                        "P=%{customdata[2]:.3e}"
                        "<extra></extra>"
                    ),
                    showlegend=False,
                )
            )

            method_fig.add_vline(
                x=1,
                line_dash="dash",
                line_color="#F1C40F"
            )

            method_fig.update_xaxes(
                type="log",
                title="OR（对数坐标）"
            )

            method_fig.update_layout(
                title=f"{selected_gene}：多种MR方法比较",
                height=500,
            )

            st.plotly_chart(
                method_fig,
                use_container_width=True
            )

            st.dataframe(
                gene_methods,
                use_container_width=True,
                hide_index=True
            )

            st.caption(
                "IVW、MR-Egger和加权中位数方法的方向越一致，"
                "结果通常越稳健；但方法一致仍不能排除全部偏倚。"
            )

    # --------------------------------------------------
    # 敏感性检查
    # --------------------------------------------------
    with sensitivity_tab:
        st.subheader("异质性、水平多效性与MR-PRESSO")

        if sensitivity is None:
            st.warning("未找到MR敏感性分析结果。")
        else:
            sensitivity = sensitivity.copy()

            for column in [
                "heterogeneity_p_value",
                "pleiotropy_p_value",
                "egger_intercept",
            ]:
                if column in sensitivity.columns:
                    sensitivity[column] = pd.to_numeric(
                        sensitivity[column],
                        errors="coerce"
                    )

            sensitivity["heterogeneity_flag"] = (
                as_boolean(
                    sensitivity[
                        "heterogeneity_flag"
                    ]
                )
            )

            sensitivity["pleiotropy_flag"] = (
                as_boolean(
                    sensitivity[
                        "pleiotropy_flag"
                    ]
                )
            )

            heterogeneity_count = int(
                sensitivity[
                    "heterogeneity_flag"
                ].sum()
            )

            pleiotropy_risk_count = int(
                sensitivity[
                    "pleiotropy_flag"
                ].sum()
            )

            presso_count = 0

            if (
                presso is not None
                and "global_test_significant"
                in presso.columns
            ):
                presso = presso.copy()

                presso[
                    "global_test_significant"
                ] = as_boolean(
                    presso[
                        "global_test_significant"
                    ]
                )

                presso_count = int(
                    presso[
                        "global_test_significant"
                    ].sum()
                )

            check1, check2, check3 = st.columns(3)

            check1.metric(
                "异质性提示",
                heterogeneity_count
            )

            check2.metric(
                "Egger多效性提示",
                pleiotropy_risk_count
            )

            check3.metric(
                "MR-PRESSO全局提示",
                presso_count
            )

            sensitivity_long = sensitivity.melt(
                id_vars=["gene_symbol"],
                value_vars=[
                    "heterogeneity_p_value",
                    "pleiotropy_p_value",
                ],
                var_name="test",
                value_name="p_value"
            )

            sensitivity_long = sensitivity_long.dropna(
                subset=["p_value"]
            )

            sensitivity_long["minus_log10_p"] = (
                -np.log10(
                    sensitivity_long[
                        "p_value"
                    ].clip(lower=1e-300)
                )
            )

            sensitivity_long["test"] = (
                sensitivity_long["test"]
                .replace(
                    {
                        "heterogeneity_p_value":
                            "IVW异质性检验",
                        "pleiotropy_p_value":
                            "MR-Egger截距检验",
                    }
                )
            )

            sensitivity_fig = px.bar(
                sensitivity_long,
                x="gene_symbol",
                y="minus_log10_p",
                color="test",
                barmode="group",
                labels={
                    "gene_symbol": "基因",
                    "minus_log10_p": "-log10(P值)",
                    "test": "敏感性检验",
                },
                title="MR敏感性检验结果",
                height=550,
            )

            sensitivity_fig.add_hline(
                y=-np.log10(0.05),
                line_dash="dash",
                line_color="#F1C40F",
                annotation_text="P = 0.05"
            )

            st.plotly_chart(
                sensitivity_fig,
                use_container_width=True
            )

            sensitivity_view = sensitivity.copy()

            if presso is not None:
                presso_columns = [
                    column
                    for column in [
                        "gene_symbol",
                        "global_test_p_value",
                        "outlier_count",
                        "global_test_significant",
                        "presso_interpretation",
                    ]
                    if column in presso.columns
                ]

                sensitivity_view = (
                    sensitivity_view.merge(
                        presso[presso_columns],
                        on="gene_symbol",
                        how="left"
                    )
                )

            st.dataframe(
                sensitivity_view,
                use_container_width=True,
                hide_index=True
            )

            sensitivity_csv = (
                sensitivity_view.to_csv(
                    index=False
                ).encode("utf-8-sig")
            )

            st.download_button(
                "下载MR敏感性结果（CSV）",
                data=sensitivity_csv,
                file_name="mr_sensitivity_results.csv",
                mime="text/csv",
                key="download_mr_sensitivity"
            )

            if evidence is not None:
                with st.expander("查看综合MR证据分级"):
                    st.dataframe(
                        evidence,
                        use_container_width=True,
                        hide_index=True
                    )

            st.error(
                "当前6个基因中有5个出现MR-Egger水平多效性信号。"
                "这些结果不能被简单表述为已证实的因果关系，"
                "应结合共定位分析、更严格的工具变量筛选、"
                "Steiger方向性检验以及独立GWAS重复验证。"
            )

    st.warning(
        "孟德尔随机化依赖工具变量相关性、独立性和排除限制假设。"
        "本模块展示的是遗传因果线索，而不是五积散治疗作用的"
        "直接临床证据，也不能证明调控某个基因必然产生治疗效果。"
    )

    # ---------------------------------------------------------------------------
# 页面：免疫浸润
# ---------------------------------------------------------------------------
def page_immune_infiltration(cfg, project_id):
    import numpy as np
    import plotly.express as px

    st.header("免疫微环境证据")
    st.info("本页解读既有项目结果。新数据请进入「进阶证据 → 免疫微环境」实际操作区，上传表达与签名运行ssGSEA；不执行去卷积或细胞比例估计。")
    st.caption("研究链条位置：多组学证据扩展，不是网络药理学的必经步骤")
    with st.expander("为什么分析免疫微环境？", expanded=True):
        st.markdown(
            """
溃疡性结肠炎与免疫失衡密切相关。本模块用于观察疾病组与正常组的
**免疫细胞相对富集模式**，并检验候选基因表达是否与这些模式相关，
从而为网络药理学候选机制增加一个表达层面的支持证据。

- **输入**：bulk 转录组表达矩阵、样本分组、免疫细胞基因集；
- **方法**：ssGSEA 相对富集评分、组间比较、相关性和 PCA；
- **输出**：差异免疫细胞、样本免疫图谱、候选基因—免疫关联；
- **不能说明**：真实细胞比例、细胞之间的通讯方向或药物治疗因果效应。

细胞通讯是另一类分析，通常需要单细胞数据及配体—受体数据库，
不能根据本页结果直接推断。
            """
        )

    files = cfg.get("data_files", {})

    def get_immune_data(file_key):
        return load_csv(
            project_path(
                project_id,
                files.get(file_key, "")
            )
        )

    scores = get_immune_data("immune_scores")
    long_scores = get_immune_data(
        "immune_scores_long"
    )
    statistics = get_immune_data(
        "immune_group_statistics"
    )
    cell_correlations = get_immune_data(
        "immune_cell_correlations"
    )
    pca = get_immune_data("immune_pca")
    gene_correlations = get_immune_data(
        "immune_gene_correlations"
    )

    if scores is None or statistics is None:
        st.info(
            "未找到免疫浸润结果，请检查project.yaml配置。"
        )
        return

    scores = scores.copy()
    statistics = statistics.copy()

    def as_boolean(series):
        if series.dtype == bool:
            return series.fillna(False)

        return (
            series
            .astype(str)
            .str.strip()
            .str.lower()
            .isin(["true", "1", "yes"])
        )

    statistics["p_adjust"] = pd.to_numeric(
        statistics["p_adjust"],
        errors="coerce"
    )

    statistics["cohen_d"] = pd.to_numeric(
        statistics["cohen_d"],
        errors="coerce"
    )

    statistics["fdr_significant"] = (
        as_boolean(
            statistics["fdr_significant"]
        )
    )

    sample_count = len(scores)

    control_count = int(
        (
            scores["group"] == "Control"
        ).sum()
    )

    uc_count = int(
        (
            scores["group"] == "UC"
        ).sum()
    )

    immune_cell_count = len(statistics)

    significant_cell_count = int(
        statistics[
            "fdr_significant"
        ].sum()
    )

    metric1, metric2, metric3, metric4 = st.columns(4)

    metric1.metric(
        "分析样本",
        sample_count,
        help=(
            f"正常对照 {control_count}，"
            f"UC {uc_count}"
        )
    )

    metric2.metric(
        "免疫细胞类型",
        immune_cell_count
    )

    metric3.metric(
        "FDR显著差异",
        significant_cell_count
    )

    metric4.metric(
        "候选基因",
        (
            gene_correlations[
                "gene_symbol"
            ].nunique()
            if gene_correlations is not None
            else 0
        )
    )

    st.caption(
        "本模块基于GSE38713表达矩阵进行ssGSEA免疫浸润估计，"
        "比较活动期溃疡性结肠炎患者与正常对照。"
        "ssGSEA得分是相对富集分数，不等同于真实细胞比例。"
    )

    overview_tab, difference_tab, correlation_tab, gene_tab = (
        st.tabs(
            [
                "样本免疫概览",
                "差异免疫细胞",
                "免疫细胞相关性",
                "候选基因关联",
            ]
        )
    )

    # --------------------------------------------------
    # 样本免疫概览
    # --------------------------------------------------
    with overview_tab:
        st.subheader("样本免疫特征")

        if pca is None:
            st.warning("未找到免疫PCA结果。")
        else:
            pca = pca.copy()

            pca["PC1"] = pd.to_numeric(
                pca["PC1"],
                errors="coerce"
            )

            pca["PC2"] = pd.to_numeric(
                pca["PC2"],
                errors="coerce"
            )

            pc1_percent = pd.to_numeric(
                pca["PC1_explained_percent"],
                errors="coerce"
            ).iloc[0]

            pc2_percent = pd.to_numeric(
                pca["PC2_explained_percent"],
                errors="coerce"
            ).iloc[0]

            pca_fig = px.scatter(
                pca,
                x="PC1",
                y="PC2",
                color="group",
                hover_name="sample_id",
                color_discrete_map={
                    "Control": "#2E8B57",
                    "UC": "#C0392B",
                },
                labels={
                    "PC1": (
                        f"PC1（{pc1_percent:.1f}%）"
                    ),
                    "PC2": (
                        f"PC2（{pc2_percent:.1f}%）"
                    ),
                    "group": "分组",
                },
                title="基于28种免疫细胞得分的PCA",
                height=560,
            )

            pca_fig.update_traces(
                marker={
                    "size": 12,
                    "line": {
                        "width": 1,
                        "color": "white",
                    },
                }
            )

            st.plotly_chart(
                pca_fig,
                use_container_width=True
            )

        st.subheader("样本免疫得分热图")

        immune_columns = [
            column
            for column in scores.columns
            if column not in [
                "sample_id",
                "label",
                "group",
            ]
        ]

        heatmap_data = (
            scores
            .set_index("sample_id")[
                immune_columns
            ]
            .apply(
                pd.to_numeric,
                errors="coerce"
            )
        )

        row_means = heatmap_data.mean(axis=1)
        row_stds = heatmap_data.std(axis=1)

        row_stds = row_stds.replace(0, 1)

        standardized_heatmap = (
            heatmap_data
            .sub(row_means, axis=0)
            .div(row_stds, axis=0)
        )

        heatmap_fig = px.imshow(
            standardized_heatmap.T,
            color_continuous_scale="RdBu_r",
            color_continuous_midpoint=0,
            aspect="auto",
            labels={
                "x": "样本",
                "y": "免疫细胞",
                "color": "样本内标准化得分",
            },
            title="免疫浸润得分热图",
            height=850,
        )

        st.plotly_chart(
            heatmap_fig,
            use_container_width=True
        )

        st.info(
            "PCA和热图用于观察样本整体免疫模式，"
            "不能单独证明某种免疫细胞参与五积散的治疗作用。"
        )

    # --------------------------------------------------
    # 差异免疫细胞
    # --------------------------------------------------
    with difference_tab:
        st.subheader("UC与正常对照的免疫差异")

        filter1, filter2 = st.columns(2)

        fdr_limit = filter1.number_input(
            "校正后P值上限",
            min_value=0.0,
            max_value=1.0,
            value=0.05,
            step=0.001,
            format="%.3f",
            key="immune_fdr_limit"
        )

        direction_option = filter2.selectbox(
            "变化方向",
            [
                "全部",
                "Higher in UC",
                "Lower in UC",
            ],
            format_func=lambda value: {
                "全部": "全部",
                "Higher in UC": "UC中升高",
                "Lower in UC": "UC中降低",
            }.get(value, value),
            key="immune_direction"
        )

        filtered_statistics = statistics[
            statistics["p_adjust"] <= fdr_limit
        ].copy()

        if direction_option != "全部":
            filtered_statistics = (
                filtered_statistics[
                    filtered_statistics["direction"]
                    == direction_option
                ]
            )

        filtered_statistics = (
            filtered_statistics
            .sort_values("cohen_d")
        )

        higher_count = int(
            (
                filtered_statistics["direction"]
                == "Higher in UC"
            ).sum()
        )

        lower_count = int(
            (
                filtered_statistics["direction"]
                == "Lower in UC"
            ).sum()
        )

        result1, result2, result3 = st.columns(3)

        result1.metric(
            "当前差异细胞",
            len(filtered_statistics)
        )

        result2.metric(
            "UC中升高",
            higher_count
        )

        result3.metric(
            "UC中降低",
            lower_count
        )

        if not filtered_statistics.empty:
            effect_fig = px.bar(
                filtered_statistics,
                x="cohen_d",
                y="immune_cell",
                orientation="h",
                color="direction",
                color_discrete_map={
                    "Higher in UC": "#C0392B",
                    "Lower in UC": "#2874A6",
                },
                hover_data={
                    "control_mean": ":.3f",
                    "uc_mean": ":.3f",
                    "p_adjust": ":.3e",
                    "cohen_d": ":.3f",
                },
                labels={
                    "cohen_d":
                        "Cohen's d（UC - Control）",
                    "immune_cell": "免疫细胞",
                    "direction": "变化方向",
                    "control_mean":
                        "正常对照均值",
                    "uc_mean": "UC均值",
                    "p_adjust":
                        "校正后P值",
                },
                title="差异免疫细胞效应量",
                height=max(
                    550,
                    len(filtered_statistics) * 32
                ),
            )

            effect_fig.add_vline(
                x=0,
                line_dash="dash",
                line_color="#F1C40F"
            )

            st.plotly_chart(
                effect_fig,
                use_container_width=True
            )
        else:
            st.warning(
                "当前条件下没有差异免疫细胞。"
            )

        st.subheader("单个免疫细胞组间分布")

        available_cells = sorted(
            statistics["immune_cell"]
            .dropna()
            .astype(str)
            .unique()
            .tolist()
        )

        selected_cell = st.selectbox(
            "选择免疫细胞",
            available_cells,
            key="immune_box_cell"
        )

        if long_scores is not None:
            selected_scores = long_scores[
                long_scores["immune_cell"]
                == selected_cell
            ].copy()

            selected_scores["ssgsea_score"] = (
                pd.to_numeric(
                    selected_scores["ssgsea_score"],
                    errors="coerce"
                )
            )

            box_fig = px.box(
                selected_scores,
                x="group",
                y="ssgsea_score",
                color="group",
                points="all",
                color_discrete_map={
                    "Control": "#2E8B57",
                    "UC": "#C0392B",
                },
                labels={
                    "group": "分组",
                    "ssgsea_score":
                        "ssGSEA得分",
                },
                title=(
                    f"{selected_cell}组间分布"
                ),
                height=520,
            )

            st.plotly_chart(
                box_fig,
                use_container_width=True
            )

        st.dataframe(
            filtered_statistics,
            use_container_width=True,
            hide_index=True
        )

        immune_csv = filtered_statistics.to_csv(
            index=False
        ).encode("utf-8-sig")

        st.download_button(
            "下载当前差异免疫结果（CSV）",
            data=immune_csv,
            file_name=(
                "immune_group_statistics.csv"
            ),
            mime="text/csv",
            key="download_immune_statistics"
        )

    # --------------------------------------------------
    # 免疫细胞相关性
    # --------------------------------------------------
    with correlation_tab:
        st.subheader("免疫细胞得分相关矩阵")

        if cell_correlations is None:
            st.warning("未找到免疫细胞相关矩阵。")
        else:
            cell_correlations = (
                cell_correlations.copy()
            )

            index_column = (
                "immune_cell"
                if "immune_cell"
                in cell_correlations.columns
                else cell_correlations.columns[0]
            )

            matrix = (
                cell_correlations
                .set_index(index_column)
                .apply(
                    pd.to_numeric,
                    errors="coerce"
                )
            )

            correlation_fig = px.imshow(
                matrix,
                color_continuous_scale="RdBu_r",
                color_continuous_midpoint=0,
                zmin=-1,
                zmax=1,
                aspect="auto",
                labels={
                    "x": "免疫细胞",
                    "y": "免疫细胞",
                    "color": "Spearman相关系数",
                },
                title="免疫细胞Spearman相关性",
                height=850,
            )

            st.plotly_chart(
                correlation_fig,
                use_container_width=True
            )

            with st.expander(
                "查看免疫细胞相关矩阵"
            ):
                st.dataframe(
                    cell_correlations,
                    use_container_width=True,
                    hide_index=True
                )

        st.caption(
            "免疫细胞间相关性表示不同ssGSEA得分在样本中的"
            "共同变化，不代表细胞之间存在直接调控或因果关系。"
        )

    # --------------------------------------------------
    # 候选基因关联
    # --------------------------------------------------
    with gene_tab:
        st.subheader(
            "候选基因表达与免疫浸润得分相关性"
        )

        if (
            gene_correlations is None
            or gene_correlations.empty
        ):
            st.warning(
                "未找到候选基因—免疫细胞相关结果。"
            )
        else:
            gene_correlations = (
                gene_correlations.copy()
            )

            for column in [
                "spearman_r",
                "p_value",
                "adjusted_p_value",
            ]:
                gene_correlations[column] = (
                    pd.to_numeric(
                        gene_correlations[column],
                        errors="coerce"
                    )
                )

            gene_options = sorted(
                gene_correlations["gene_symbol"]
                .dropna()
                .astype(str)
                .unique()
                .tolist()
            )

            selected_gene = st.selectbox(
                "选择候选基因",
                gene_options,
                key="immune_gene"
            )

            only_significant = st.checkbox(
                "只显示FDR显著相关",
                value=True,
                key="immune_gene_significant"
            )

            selected_correlations = (
                gene_correlations[
                    gene_correlations["gene_symbol"]
                    == selected_gene
                ].copy()
            )

            if only_significant:
                selected_correlations = (
                    selected_correlations[
                        selected_correlations[
                            "adjusted_p_value"
                        ] < 0.05
                    ]
                )

            selected_correlations[
                "absolute_correlation"
            ] = (
                selected_correlations[
                    "spearman_r"
                ].abs()
            )

            selected_correlations = (
                selected_correlations
                .sort_values(
                    "absolute_correlation",
                    ascending=False
                )
            )

            plot_correlations = (
                selected_correlations
                .head(20)
                .sort_values("spearman_r")
            )

            if not plot_correlations.empty:
                gene_fig = px.bar(
                    plot_correlations,
                    x="spearman_r",
                    y="immune_cell",
                    orientation="h",
                    color="spearman_r",
                    color_continuous_scale="RdBu_r",
                    color_continuous_midpoint=0,
                    range_color=[-1, 1],
                    hover_data={
                        "adjusted_p_value": ":.3e",
                        "sample_count": True,
                    },
                    labels={
                        "spearman_r":
                            "Spearman相关系数",
                        "immune_cell":
                            "免疫细胞",
                        "adjusted_p_value":
                            "FDR",
                    },
                    title=(
                        f"{selected_gene}与免疫细胞"
                        "相关性 Top 20"
                    ),
                    height=650,
                )

                gene_fig.add_vline(
                    x=0,
                    line_dash="dash",
                    line_color="#F1C40F"
                )

                st.plotly_chart(
                    gene_fig,
                    use_container_width=True
                )
            else:
                st.warning(
                    "当前筛选条件下没有相关结果。"
                )

            st.dataframe(
                selected_correlations,
                use_container_width=True,
                hide_index=True
            )

            gene_csv = (
                selected_correlations.to_csv(
                    index=False
                ).encode("utf-8-sig")
            )

            st.download_button(
                "下载当前基因免疫相关结果（CSV）",
                data=gene_csv,
                file_name=(
                    f"{selected_gene}_immune_correlations.csv"
                ),
                mime="text/csv",
                key="download_immune_gene"
            )

        st.error(
            "基因表达和ssGSEA得分来自同一表达矩阵，"
            "因此相关性可能包含数学耦合和循环分析效应。"
            "这些结果只能用于提出假设，不能作为独立验证、"
            "细胞比例证据或治疗因果证据。"
        )

# ---------------------------------------------------------------------------
# 页面：综合证据
# ---------------------------------------------------------------------------
def page_integrated_evidence(cfg, project_id):
    import plotly.express as px

    st.header("综合证据整合")

    files = cfg.get("data_files", {})

    matrix = load_csv(
        project_path(
            project_id,
            files.get("integrated_evidence_matrix", ""),
        )
    )
    summary = load_csv(
        project_path(
            project_id,
            files.get("integrated_evidence_summary", ""),
        )
    )

    if matrix is None:
        st.info(
            "未找到综合证据矩阵，请先运行"
            " prepare_integrated_evidence.py。"
        )
        return

    matrix = matrix.copy()

    def as_boolean(series):
        if series.dtype == bool:
            return series.fillna(False)

        return (
            series.astype(str)
            .str.strip()
            .str.lower()
            .isin(["true", "1", "yes"])
        )

    boolean_columns = [
        "network_supported",
        "ppi_top20_hub",
        "transcriptome_supported",
        "ml_supported",
        "mr_analyzed",
        "mr_fdr_supported",
        "immune_supported",
    ]

    for column in boolean_columns:
        if column in matrix.columns:
            matrix[column] = as_boolean(matrix[column])

    numeric_columns = [
        "evidence_layer_count",
        "degree",
        "logFC",
        "transcriptome_fdr",
        "ml_model_count",
        "ml_mean_importance",
        "mr_odds_ratio",
        "mr_fdr",
        "immune_significant_cells",
        "immune_max_abs_r",
        "immune_top_r",
    ]

    for column in numeric_columns:
        if column in matrix.columns:
            matrix[column] = pd.to_numeric(
                matrix[column],
                errors="coerce",
            )

    total_genes = len(matrix)
    priority_a = int(
        (
            matrix["priority_tier"]
            == "A：多模块重点候选"
        ).sum()
    )
    three_layer_genes = int(
        (
            matrix["evidence_layer_count"] >= 3
        ).sum()
    )
    mr_supported = int(
        matrix["mr_fdr_supported"].sum()
    )
    immune_supported = int(
        matrix["immune_supported"].sum()
    )

    card1, card2, card3, card4, card5 = st.columns(5)

    card1.metric("整合基因", f"{total_genes:,}")
    card2.metric("A级重点候选", f"{priority_a:,}")
    card3.metric("≥3层证据", f"{three_layer_genes:,}")
    card4.metric("MR-FDR显著", f"{mr_supported:,}")
    card5.metric("免疫相关支持", f"{immune_supported:,}")

    st.caption(
        "本页面将网络药理学、PPI拓扑、转录组、机器学习、"
        "孟德尔随机化和免疫相关结果按基因统一汇总。"
    )

    overview_tab, genes_tab, detail_tab, note_tab = st.tabs(
        [
            "证据概览",
            "重点候选基因",
            "单基因证据卡",
            "方法说明",
        ]
    )

    with overview_tab:
        st.subheader("多模块证据分布")

        left, right = st.columns(2)

        tier_order = [
            "A：多模块重点候选",
            "B：较强候选",
            "C：初步候选",
            "D：单一来源",
        ]

        tier_counts = (
            matrix["priority_tier"]
            .value_counts()
            .reindex(tier_order, fill_value=0)
            .rename_axis("优先级")
            .reset_index(name="基因数")
        )

        tier_figure = px.bar(
            tier_counts,
            x="优先级",
            y="基因数",
            color="优先级",
            text="基因数",
            title="候选基因优先级分布",
            color_discrete_map={
                "A：多模块重点候选": "#C0392B",
                "B：较强候选": "#E67E22",
                "C：初步候选": "#F1C40F",
                "D：单一来源": "#5D6D7E",
            },
        )
        tier_figure.update_layout(
            showlegend=False,
            xaxis_title="",
        )
        left.plotly_chart(
            tier_figure,
            use_container_width=True,
        )

        source_labels = {
            "network_supported": "网络药理学",
            "transcriptome_supported": "转录组",
            "ml_supported": "机器学习",
            "mr_fdr_supported": "MR-FDR显著",
            "immune_supported": "免疫相关",
        }

        source_data = pd.DataFrame(
            [
                {
                    "证据模块": label,
                    "支持基因数": int(matrix[column].sum()),
                }
                for column, label in source_labels.items()
            ]
        )

        source_figure = px.bar(
            source_data,
            x="证据模块",
            y="支持基因数",
            color="支持基因数",
            text="支持基因数",
            title="各分析模块支持基因数",
            color_continuous_scale="Tealgrn",
        )
        source_figure.update_layout(
            coloraxis_showscale=False,
            xaxis_title="",
        )
        right.plotly_chart(
            source_figure,
            use_container_width=True,
        )

        layer_counts = (
            matrix["evidence_layer_count"]
            .value_counts()
            .sort_index()
            .rename_axis("证据层数")
            .reset_index(name="基因数")
        )

        layer_figure = px.bar(
            layer_counts,
            x="证据层数",
            y="基因数",
            text="基因数",
            title="综合证据层数分布",
            color="证据层数",
            color_continuous_scale="Reds",
        )
        layer_figure.update_layout(
            coloraxis_showscale=False,
        )
        st.plotly_chart(
            layer_figure,
            use_container_width=True,
        )

    with genes_tab:
        st.subheader("重点候选基因筛选")

        filter1, filter2, filter3 = st.columns(3)

        minimum_layers = filter1.slider(
            "最低证据层数",
            min_value=1,
            max_value=5,
            value=3,
        )

        tier_options = [
            "全部",
            *[
                value
                for value in tier_order
                if value in matrix["priority_tier"].unique()
            ],
        ]

        selected_tier = filter2.selectbox(
            "候选优先级",
            tier_options,
        )

        require_mr = filter3.checkbox(
            "仅显示MR-FDR显著",
            value=False,
        )

        filtered = matrix[
            matrix["evidence_layer_count"]
            >= minimum_layers
        ].copy()

        if selected_tier != "全部":
            filtered = filtered[
                filtered["priority_tier"]
                == selected_tier
            ]

        if require_mr:
            filtered = filtered[
                filtered["mr_fdr_supported"]
            ]

        st.metric(
            "当前筛选基因数",
            len(filtered),
        )

        table_columns = [
            "priority_rank",
            "gene_symbol",
            "priority_tier",
            "evidence_layer_count",
            "evidence_sources",
            "degree",
            "logFC",
            "transcriptome_fdr",
            "ml_model_count",
            "mr_odds_ratio",
            "mr_fdr",
            "immune_significant_cells",
            "immune_top_cell",
            "immune_top_r",
        ]
        table_columns = [
            column
            for column in table_columns
            if column in filtered.columns
        ]

        st.dataframe(
            filtered[table_columns],
            use_container_width=True,
            hide_index=True,
        )

        st.download_button(
            "下载当前筛选结果",
            data=filtered.to_csv(
                index=False,
                encoding="utf-8-sig",
            ),
            file_name=(
                f"{project_id}_integrated_candidates.csv"
            ),
            mime="text/csv",
        )

    with detail_tab:
        st.subheader("单基因综合证据卡")

        selected_gene = st.selectbox(
            "选择候选基因",
            matrix["gene_symbol"].tolist(),
        )

        gene_row = matrix[
            matrix["gene_symbol"] == selected_gene
        ].iloc[0]

        detail1, detail2, detail3, detail4 = st.columns(4)

        detail1.metric(
            "综合排名",
            int(gene_row["priority_rank"]),
        )
        detail2.metric(
            "证据层数",
            int(gene_row["evidence_layer_count"]),
        )
        detail3.metric(
            "PPI Degree",
            int(gene_row.get("degree", 0) or 0),
        )
        detail4.metric(
            "显著相关免疫细胞",
            int(
                gene_row.get(
                    "immune_significant_cells",
                    0,
                )
                or 0
            ),
        )

        st.write(
            f"**优先级：** "
            f"{gene_row['priority_tier']}"
        )
        st.write(
            f"**证据来源：** "
            f"{gene_row['evidence_sources']}"
        )

        evidence_card = pd.DataFrame(
            [
                {
                    "证据模块": "网络药理学",
                    "是否支持": (
                        "是"
                        if gene_row["network_supported"]
                        else "否"
                    ),
                    "主要结果": "复方—疾病共同靶点",
                },
                {
                    "证据模块": "PPI网络",
                    "是否支持": (
                        "Top 20枢纽"
                        if gene_row["ppi_top20_hub"]
                        else "已分析"
                    ),
                    "主要结果": (
                        f"Degree = "
                        f"{int(gene_row.get('degree', 0) or 0)}"
                    ),
                },
                {
                    "证据模块": "转录组",
                    "是否支持": (
                        "是"
                        if gene_row[
                            "transcriptome_supported"
                        ]
                        else "否"
                    ),
                    "主要结果": (
                        f"logFC = "
                        f"{gene_row.get('logFC', float('nan')):.3f}"
                    ),
                },
                {
                    "证据模块": "机器学习",
                    "是否支持": (
                        "是"
                        if gene_row["ml_supported"]
                        else "否"
                    ),
                    "主要结果": (
                        f"{int(gene_row.get('ml_model_count', 0) or 0)}"
                        " 个模型支持"
                    ),
                },
                {
                    "证据模块": "孟德尔随机化",
                    "是否支持": (
                        "FDR显著"
                        if gene_row["mr_fdr_supported"]
                        else (
                            "已分析但未达FDR显著"
                            if gene_row["mr_analyzed"]
                            else "未分析"
                        )
                    ),
                    "主要结果": (
                        f"OR = "
                        f"{gene_row.get('mr_odds_ratio', float('nan')):.3f}"
                    ),
                },
                {
                    "证据模块": "免疫相关",
                    "是否支持": (
                        "是"
                        if gene_row["immune_supported"]
                        else "否"
                    ),
                    "主要结果": (
                        f"最强相关细胞："
                        f"{gene_row.get('immune_top_cell', '-')}"
                    ),
                },
            ]
        )

        st.dataframe(
            evidence_card,
            use_container_width=True,
            hide_index=True,
        )

    with note_tab:
        st.subheader("综合证据的正确解释")

        st.info(
            "证据层数表示一个基因被多少类分析模块支持，"
            "用于结果整理和候选优先级排序。"
        )

        st.warning(
            "转录组、机器学习和免疫浸润部分可能使用同一表达"
            "数据集，因此不能被视为完全独立的重复验证。"
        )

        st.warning(
            "PPI中心性表示网络拓扑位置，不等同于疾病因果性；"
            "MR结果还需要结合工具变量强度、异质性和多效性判断。"
        )

        st.error(
            "综合排名用于科研假设生成，不构成临床诊断、"
            "治疗推荐或五积散疗效证明。"
        )

        if summary is not None:
            with st.expander("查看综合分析汇总记录"):
                st.dataframe(
                    summary,
                    use_container_width=True,
                    hide_index=True,
                )        

def page_generic_table(cfg, project_id, title, file_key):
    st.header(title)
    files = cfg.get("data_files", {})
    df = load_csv(project_path(project_id, files.get(file_key, "")))
    if df is None:
        st.info(f"未找到 {file_key} 对应数据文件。")
        return
    st.dataframe(df, use_container_width=True, hide_index=True)


# ---------------------------------------------------------------------------
# 页面：报告中心
# ---------------------------------------------------------------------------
def page_report(cfg, project_id):
    import io
    import zipfile
    from datetime import datetime

    from docx import Document

    st.header("报告中心")
    st.caption("现有项目报告与教学自查记录用途不同：学习项目可在「项目工具 → 资料模板与交付自查」填写并下载交付记录；平台不会自动核验外部计算结果。")
    st.write(
        "汇总项目结论，导出 Word 报告、Excel 工作簿"
        "以及完整 CSV 数据包。"
    )

    files = cfg.get("data_files", {})
    frames = {}

    for key, relative_path in files.items():
        dataframe = load_csv(
            project_path(project_id, relative_path)
        )
        if dataframe is not None:
            frames[key] = dataframe

    if not frames:
        st.info("当前项目暂无可导出的数据文件。")
        return

    def summary_value(frame_key, metric, default="-"):
        dataframe = frames.get(frame_key)

        if dataframe is None:
            return default

        if not {"metric", "value"}.issubset(
            dataframe.columns
        ):
            return default

        matched = dataframe[
            dataframe["metric"].astype(str) == metric
        ]

        if matched.empty:
            return default

        return matched.iloc[0]["value"]

    integrated = frames.get(
        "integrated_evidence_matrix"
    )

    if integrated is not None:
        priority_a_count = int(
            (
                integrated["priority_tier"]
                == "A：多模块重点候选"
            ).sum()
        )

        top_genes = (
            integrated.sort_values("priority_rank")
            .head(10)["gene_symbol"]
            .astype(str)
            .tolist()
        )
    else:
        priority_a_count = 0
        top_genes = []

    total_records = sum(
        len(dataframe)
        for dataframe in frames.values()
    )

    card1, card2, card3, card4 = st.columns(4)

    card1.metric("已接入数据表", len(frames))
    card2.metric("总数据记录", f"{total_records:,}")
    card3.metric("A级重点候选", priority_a_count)
    card4.metric("可导出格式", "3种")

    transcriptome_degs = summary_value(
        "transcriptome_summary",
        "significant_degs",
    )
    best_auc = summary_value(
        "ml_summary",
        "best_external_auc",
    )
    ml_features = summary_value(
        "ml_summary",
        "multi_model_features",
    )
    mr_fdr_genes = summary_value(
        "mr_summary",
        "full_snp_fdr_significant",
    )
    significant_immune_cells = summary_value(
        "immune_summary",
        "significant_immune_cells",
    )

    top_gene_text = (
        "、".join(top_genes)
        if top_genes
        else "暂无"
    )

    generated_time = datetime.now().strftime(
        "%Y-%m-%d %H:%M"
    )

    report_markdown = f"""
# {cfg.get("project_name", project_id)} 综合分析报告

生成时间：{generated_time}

## 一、项目概况

- 研究复方：{cfg.get("formula_name", "-")}
- 研究疾病：{cfg.get("disease_name", "-")}
- 研究物种：{cfg.get("species", "-")}
- 项目版本：{cfg.get("version", "-")}

## 二、主要分析结果

- 网络药理学共同靶点：{len(frames.get("target_overlap", []))} 个
- 转录组显著差异基因：{transcriptome_degs} 个
- 机器学习多模型候选特征：{ml_features} 个
- 机器学习最佳外部验证 AUC：{best_auc}
- MR 全工具变量分析中 FDR 显著基因：{mr_fdr_genes} 个
- FDR 显著差异免疫细胞：{significant_immune_cells} 类
- 综合证据 A 级重点候选：{priority_a_count} 个

## 三、综合证据重点基因

当前综合排序前 10 位基因：

{top_gene_text}

综合排序用于整合网络药理学、PPI 拓扑、转录组、
机器学习、孟德尔随机化和免疫相关分析结果。

## 四、结果解释

多模块支持说明候选基因在多个分析环节中被重复识别，
适合优先开展文献验证、独立队列验证和实验验证。

转录组、机器学习和免疫浸润可能使用同一表达数据集，
因此这些模块不能被视为完全独立的重复证据。

PPI 中心性反映网络拓扑位置，不直接代表疾病因果关系。
MR 结果还需要结合工具变量强度、异质性、水平多效性
和独立数据进行谨慎解释。

## 五、研究声明

{cfg.get("disclaimer", DISCLAIMER)}

本报告中的综合排名用于科研假设生成，不构成临床诊断、
治疗建议或五积散疗效证明。
""".strip()

    report_tab, export_tab, reproducibility_tab = st.tabs(
        [
            "项目摘要",
            "数据导出",
            "可复现性清单",
        ]
    )

    with report_tab:
        st.subheader("自动生成的项目摘要")
        st.markdown(report_markdown)

        markdown_bytes = report_markdown.encode(
            "utf-8-sig"
        )

        # 生成 Word 报告
        document = Document()

        document.add_heading(
            f"{cfg.get('project_name', project_id)}"
            " 综合分析报告",
            level=0,
        )

        document.add_paragraph(
            f"生成时间：{generated_time}"
        )

        document.add_heading("一、项目概况", level=1)
        document.add_paragraph(
            f"研究复方：{cfg.get('formula_name', '-')}"
        )
        document.add_paragraph(
            f"研究疾病：{cfg.get('disease_name', '-')}"
        )
        document.add_paragraph(
            f"研究物种：{cfg.get('species', '-')}"
        )
        document.add_paragraph(
            f"项目版本：{cfg.get('version', '-')}"
        )

        document.add_heading(
            "二、主要分析结果",
            level=1,
        )

        result_items = [
            (
                "网络药理学共同靶点",
                f"{len(frames.get('target_overlap', []))} 个",
            ),
            (
                "转录组显著差异基因",
                f"{transcriptome_degs} 个",
            ),
            (
                "机器学习多模型候选特征",
                f"{ml_features} 个",
            ),
            (
                "机器学习最佳外部验证 AUC",
                str(best_auc),
            ),
            (
                "MR中FDR显著基因",
                f"{mr_fdr_genes} 个",
            ),
            (
                "显著差异免疫细胞",
                f"{significant_immune_cells} 类",
            ),
            (
                "A级重点候选",
                f"{priority_a_count} 个",
            ),
        ]

        result_table = document.add_table(
            rows=1,
            cols=2,
        )
        result_table.style = "Table Grid"
        result_table.rows[0].cells[0].text = "分析指标"
        result_table.rows[0].cells[1].text = "结果"

        for indicator, result in result_items:
            cells = result_table.add_row().cells
            cells[0].text = indicator
            cells[1].text = result

        document.add_heading(
            "三、综合证据重点基因",
            level=1,
        )

        if integrated is not None:
            gene_columns = [
                "gene_symbol",
                "priority_tier",
                "evidence_layer_count",
                "evidence_sources",
            ]

            top_table_data = (
                integrated.sort_values("priority_rank")
                .head(10)[gene_columns]
            )

            gene_table = document.add_table(
                rows=1,
                cols=4,
            )
            gene_table.style = "Table Grid"

            headers = [
                "基因",
                "优先级",
                "证据层数",
                "证据来源",
            ]

            for index, header in enumerate(headers):
                gene_table.rows[0].cells[index].text = header

            for _, row in top_table_data.iterrows():
                cells = gene_table.add_row().cells
                cells[0].text = str(row["gene_symbol"])
                cells[1].text = str(row["priority_tier"])
                cells[2].text = str(
                    row["evidence_layer_count"]
                )
                cells[3].text = str(
                    row["evidence_sources"]
                )
        else:
            document.add_paragraph(
                "尚未生成综合证据矩阵。"
            )

        document.add_heading(
            "四、结果解释与限制",
            level=1,
        )
        document.add_paragraph(
            "多模块支持用于候选优先级排序，"
            "不代表各证据完全独立。"
        )
        document.add_paragraph(
            "转录组、机器学习和免疫浸润可能使用"
            "同一表达数据集，存在相关性和循环分析风险。"
        )
        document.add_paragraph(
            "PPI中心性不等同于因果性；MR结果需要结合"
            "工具变量强度、异质性和多效性解释。"
        )

        document.add_heading(
            "五、研究声明",
            level=1,
        )
        document.add_paragraph(
            cfg.get("disclaimer", DISCLAIMER)
        )
        document.add_paragraph(
            "综合排名用于科研假设生成，不构成临床诊断、"
            "治疗推荐或疗效证明。"
        )

        word_buffer = io.BytesIO()
        document.save(word_buffer)

        download1, download2 = st.columns(2)

        download1.download_button(
            "下载 Word 综合报告",
            data=word_buffer.getvalue(),
            file_name=(
                f"{project_id}_综合分析报告.docx"
            ),
            mime=(
                "application/vnd.openxmlformats-"
                "officedocument.wordprocessingml.document"
            ),
        )

        download2.download_button(
            "下载 Markdown 摘要",
            data=markdown_bytes,
            file_name=(
                f"{project_id}_综合分析摘要.md"
            ),
            mime="text/markdown",
        )

    with export_tab:
        st.subheader("完整分析数据导出")

        st.write(
            "可以选择适合后续统计分析、审阅或存档的格式。"
        )

        # 全部 CSV 文件压缩包
        zip_buffer = io.BytesIO()

        with zipfile.ZipFile(
            zip_buffer,
            "w",
            zipfile.ZIP_DEFLATED,
        ) as zip_file:
            for key, dataframe in frames.items():
                zip_file.writestr(
                    f"{key}.csv",
                    dataframe.to_csv(
                        index=False,
                        encoding="utf-8-sig",
                    ),
                )

            zip_file.writestr(
                "project.yaml",
                yaml.safe_dump(
                    cfg,
                    allow_unicode=True,
                    sort_keys=False,
                ),
            )

            zip_file.writestr(
                "综合分析摘要.md",
                markdown_bytes,
            )

        # 全部数据写入一个 Excel 工作簿
        excel_buffer = io.BytesIO()

        with pd.ExcelWriter(
            excel_buffer,
            engine="openpyxl",
        ) as writer:
            for index, (key, dataframe) in enumerate(
                frames.items(),
                start=1,
            ):
                sheet_name = f"{index:02d}_{key}"[:31]

                dataframe.to_excel(
                    writer,
                    sheet_name=sheet_name,
                    index=False,
                )

        export1, export2 = st.columns(2)

        export1.download_button(
            "下载全部 CSV 数据包",
            data=zip_buffer.getvalue(),
            file_name=f"{project_id}_全部数据.zip",
            mime="application/zip",
        )

        export2.download_button(
            "下载 Excel 汇总工作簿",
            data=excel_buffer.getvalue(),
            file_name=f"{project_id}_全部数据.xlsx",
            mime=(
                "application/vnd.openxmlformats-"
                "officedocument.spreadsheetml.sheet"
            ),
        )

        file_overview = pd.DataFrame(
            [
                {
                    "数据键": key,
                    "记录数": len(dataframe),
                    "字段数": len(dataframe.columns),
                    "配置路径": files.get(key, "-"),
                }
                for key, dataframe in frames.items()
            ]
        )

        st.dataframe(
            file_overview,
            use_container_width=True,
            hide_index=True,
        )

    with reproducibility_tab:
        st.subheader("项目可复现性清单")

        required_scripts = [
            "src/prepare_day1_data.py",
            "src/prepare_ppi_data.py",
            "src/prepare_transcriptomics_data.py",
            "src/prepare_machine_learning_data.py",
            "src/prepare_mr_data.py",
            "src/prepare_immune_data.py",
            "src/prepare_integrated_evidence.py",
        ]

        script_status = pd.DataFrame(
            [
                {
                    "处理程序": script,
                    "状态": (
                        "存在"
                        if os.path.isfile(
                            os.path.join(
                                BASE_DIR,
                                script,
                            )
                        )
                        else "缺失"
                    ),
                }
                for script in required_scripts
            ]
        )

        st.dataframe(
            script_status,
            use_container_width=True,
            hide_index=True,
        )

        missing_files = []

        for key, relative_path in files.items():
            if not os.path.isfile(
                project_path(
                    project_id,
                    relative_path,
                )
            ):
                missing_files.append(
                    f"{key}: {relative_path}"
                )

        if missing_files:
            st.error(
                "以下配置文件缺失：\n\n"
                + "\n".join(
                    f"- {item}"
                    for item in missing_files
                )
            )
        else:
            st.success(
                "project.yaml 中登记的数据文件均存在。"
            )

        st.info(
            "建议每次更新原始数据后，依次重新运行相应的"
            "数据处理程序，最后重新运行"
            " prepare_integrated_evidence.py。"
        )

    st.divider()
    st.info(cfg.get("disclaimer", DISCLAIMER))


# ---------------------------------------------------------------------------
# 主程序
# ---------------------------------------------------------------------------
PAGES = {
    "作品首页": page_product_home,
    "分析工作台": page_project_workspace,
    "教学实训中心": page_learning_center,
    "网络药理学操作课程": page_learning_paths,
    "网络药理学教学实验": page_ppi,
    "项目管理": page_project_management,
    "新建学习任务": page_project_creation,
    "项目首页": page_home,
    "药材-候选靶点": page_herbs,
    "疾病基因与交集": page_overlap,
    "网络药理学 / PPI": page_ppi_results,
    "GO/KEGG 教学实验": page_enrichment,
    "转录组验证": page_transcriptome,
    "机器学习结果": page_machine_learning,
    "孟德尔随机化": page_mendelian_randomization,
    "免疫微环境证据": page_immune_infiltration,
    "综合证据": page_integrated_evidence,
    "报告中心": page_report,
    "使用指南": page_user_guide,
}

NAVIGATION = {
    "教学实训": [
        "教学实训中心",
        "网络药理学操作课程",
        "网络药理学教学实验",
        "GO/KEGG 教学实验",
    ],
    "五积散案例": [
        "项目首页",
        "药材-候选靶点",
        "疾病基因与交集",
        "网络药理学 / PPI",
        "转录组验证",
        "机器学习结果",
        "孟德尔随机化",
        "免疫微环境证据",
        "综合证据",
    ],
    "项目工具": [
        "项目管理",
        "报告中心",
        "使用指南",
    ],
}


def navigate_to(route_id):
    """通过稳定URL路由导航；回调在下一次页面渲染前执行。"""
    st.query_params["view"] = resolve_route(route_id)


def open_demo_workspace():
    """Open the first project explicitly marked as a competition demo."""
    for name in list_projects():
        if load_project_config(name).get("project_type") == "competition_demo":
            st.session_state["active_project"] = name
            break
    st.query_params["view"] = "workspace"


def open_wujisan_case():
    if "wujisan_uc" in list_projects():
        st.session_state["active_project"] = "wujisan_uc"
    st.query_params["view"] = "case_home"


def main():
    st.set_page_config(**PAGE_CONFIG)
    st.markdown("""<style>
    @media (min-width: 1000px) {
      [data-testid="stSidebar"] {min-width: 17rem; max-width: 19rem; width: 19rem;}
    }
    [data-testid="stSidebarContent"] {padding-top: .5rem;}
    [data-testid="stMainBlockContainer"] {padding-top: 4.5rem;}
    [data-testid="stSidebarHeader"] {height: 2.5rem; min-height: 2.5rem;}
    [data-testid="stSidebarUserContent"] {padding-top: 0;}
    [data-testid="stSidebar"] [data-testid="stVerticalBlock"] {gap: .65rem;}
    [data-testid="stSidebar"] [data-testid="stButton"] button {
      min-height: 2.15rem; text-align: left; justify-content: flex-start;
    }
    </style>""", unsafe_allow_html=True)
    projects = list_projects()
    if not projects:
        st.error("未找到任何项目。请先在 projects/ 下创建项目目录并放置 project.yaml。")
        st.stop()

    active = resolve_route(st.query_params.get("view"))
    group, label, page, section = ROUTES[active]
    st.sidebar.markdown("### 🌿 HerbInsight")
    st.sidebar.caption("按任务进入 · 无需切换使用模式")
    with st.sidebar.expander("项目与设置", expanded=False):
        preferred_project = "wujisan_uc" if "wujisan_uc" in projects else projects[0]
        if st.session_state.get("active_project") not in projects:
            st.session_state["active_project"] = preferred_project
        from src.study_progress import render_project_controls
        render_project_controls()
        project_id = st.selectbox("数据项目", projects, key="active_project")
        st.caption("数据项目决定工作台与案例页面读取的文件；教学课程结构保持不变。")
    cfg = load_project_config(project_id)
    st.sidebar.caption("数据项目：" + str(cfg.get("project_name", project_id)))
    for group_name, entries in NAV_GROUPS.items():
        with st.sidebar.expander(group_name, expanded=group_name == group):
            for route_id, title, _, _ in entries:
                st.button(title, key="nav_" + route_id, type="primary" if active == route_id else "secondary",
                          width="stretch", on_click=navigate_to, args=(route_id,))
    with st.sidebar.expander("使用声明"):
        st.caption(DISCLAIMER)
    breadcrumb, previous_col, next_col = st.columns([3, 1.8, 1.8])
    breadcrumb.caption(group + " / " + label)
    previous, next_route = course_neighbors(active)
    if previous:
        previous_col.button("← 上一节", key="course_previous", on_click=navigate_to, args=(previous,), width="stretch")
    if next_route:
        next_col.button("下一节 →", key="course_next", on_click=navigate_to, args=(next_route,), width="stretch")
    if section is not None:
        practice = None
        if section == "ppi":
            practice = lambda: page_network_lab(cfg, project_id, embedded=True)
        elif section == "enrichment":
            practice = lambda: page_enrichment_lab(cfg, project_id, embedded=True)
        page_learning_paths(cfg, project_id, section=section, practice=practice)
    else:
        PAGES[page](cfg, project_id)


def vcfg(cfg):
    return f"版本 {cfg.get('version', '-')}"


if __name__ == "__main__":
    main()
